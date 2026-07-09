from openpyxl.styles import PatternFill


def paint_rows_by_dict(ws, data_dict: dict, fill: PatternFill, columns=None):
    for (row, _fio_norm, fio_value, pos_or_eararchy_value) in data_dict.keys():

        # 🔹 одна колонка
        if isinstance(columns, int):
            ws.cell(row=row, column=columns).fill = fill

        # 🔹 несколько колонок
        elif isinstance(columns, (list, tuple)):
            for col in columns:
                ws.cell(row=row, column=col).fill = fill

        # 🔹 вся строка
        else:
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).fill = fill
