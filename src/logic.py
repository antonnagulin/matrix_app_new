"""Оркестрация проверки сотрудников (run_check) и реэкспорт публичного API."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from .check_change_fio import (
    build_tab_index,
    change_fio,
    prepare_rename_df,
    resolve_via_rename,
)
from .check_dialogs import SKIP_ALL_DUPLICATES, ask_include_missing, choose_actual_record
from .check_duplicate import check_duplicate, is_joint, second_check_duplicate
from .constants import (
    FIO_COL,
    GREEN_FILL_COLOR,
    ORANGE_FILL_COLOR,
    RED_FILL_COLOR,
    YELLOW_FILL_COLOR,
)
from .file_utils import (
    find_header_row,
    get_fill,
    iter_column_values,
    load_check_file,
    read_sprav_file,
    row_values,
    save_file,
)
from .filters import is_dismissed, join_with_slash, normalize_fio
from .fio_index import (
    build_fio_index,
    full_fio_key,
    lookup_fio,
    person_group_key,
    pick_default_duplicate_record,
)
from .paint_row_for_dict import paint_rows_by_dict
from .results_workbook import (
    COL_COMMENT_NAME,
    COL_DUP_MANUAL_NAME,
    COL_DUP_STEP2_NAME,
    COL_DUP_STEP3_NAME,
    COL_HIERARCHY_NAME,
    COL_POSITION_NAME,
    COMMENT_STATUS,
    FINAL_EXCLUDE_COLUMNS,
    MSG_DISMISSED,
    MSG_DUP_ALL_DISMISSED,
    MSG_DUP_DEFAULT,
    MSG_DUP_JOINTS_ONLY,
    MSG_DUP_MANUAL_ACTIVE,
    MSG_DUP_MANUAL_DISMISSED,
    MSG_DUP_STEP1,
    MSG_DUP_STEP1_NO_CHECK,
    MSG_DUP_STEP2_GREEN,
    MSG_DUP_STEP2_YELLOW,
    MSG_DUP_STEP3_GREEN,
    MSG_DUP_STEP3_YELLOW,
    MSG_DUP_UNRESOLVED_INCLUDED,
    MSG_NOT_FOUND,
    MSG_RENAMED_FOUND,
    STATUS_GREEN,
    STATUS_ORANGE,
    STATUS_RED,
    STATUS_YELLOW,
    UNRESOLVED_DUP_COMMENTS,
    build_final_workbook,
    collect_highlighted_rows,
    is_accepted_for_final,
    status_from_comment,
    summarize_check_rows,
)
from .validation import validate_output_not_input, validate_sprav_schema
from .write_position_hierarchy import write_position_hierarchy

# Реэкспорт для тестов / UI (стабильный публичный API)
__all__ = [
    "SKIP_ALL_DUPLICATES",
    "ask_include_missing",
    "choose_actual_record",
    "run_check",
    "build_final_workbook",
    "collect_highlighted_rows",
    "is_accepted_for_final",
    "status_from_comment",
    "summarize_check_rows",
    "MSG_NOT_FOUND",
    "MSG_DISMISSED",
    "MSG_DUP_STEP1",
    "MSG_DUP_STEP1_NO_CHECK",
    "MSG_DUP_DEFAULT",
    "MSG_DUP_STEP2_YELLOW",
    "MSG_DUP_STEP2_GREEN",
    "MSG_DUP_STEP3_YELLOW",
    "MSG_DUP_STEP3_GREEN",
    "MSG_DUP_ALL_DISMISSED",
    "MSG_DUP_JOINTS_ONLY",
    "MSG_DUP_MANUAL_ACTIVE",
    "MSG_DUP_MANUAL_DISMISSED",
    "MSG_RENAMED_FOUND",
    "MSG_DUP_UNRESOLVED_INCLUDED",
    "STATUS_RED",
    "STATUS_ORANGE",
    "STATUS_YELLOW",
    "STATUS_GREEN",
    "COMMENT_STATUS",
    "UNRESOLVED_DUP_COMMENTS",
    "COL_HIERARCHY_NAME",
    "COL_POSITION_NAME",
    "COL_COMMENT_NAME",
    "COL_DUP_STEP2_NAME",
    "COL_DUP_STEP3_NAME",
    "COL_DUP_MANUAL_NAME",
    "FINAL_EXCLUDE_COLUMNS",
]


def _ask_include(parent, fio: str) -> str:
    if hasattr(parent, "ask_include_missing"):
        return parent.ask_include_missing(fio)
    return ask_include_missing(parent, fio)


def _choose_duplicate(parent, fio, records):
    if hasattr(parent, "choose_actual_record"):
        return parent.choose_actual_record(fio, records)
    return choose_actual_record(parent, fio, records)


def _record_as_dict(record) -> dict:
    if hasattr(record, "to_dict"):
        return record.to_dict()
    return dict(record)


def _write_sprav_columns(ws, row: int, record, hierarchy_col: int, position_col: int) -> None:
    rec = _record_as_dict(record)
    ws.cell(row=row, column=hierarchy_col, value=join_with_slash(rec))
    ws.cell(row=row, column=position_col, value=rec.get("Штатная должность", ""))


def _write_comment(ws, row: int, comment_col: int, text: str) -> None:
    ws.cell(row=row, column=comment_col, value=text)


def _comment_rows(ws, data_dict: dict, comment_col: int, text: str) -> None:
    for (row, *_rest) in data_dict.keys():
        _write_comment(ws, row, comment_col, text)


def run_check(
    sprav_file,
    check_file,
    output_file,
    check_dismissed=True,
    check_duplicates=True,
    prompt_duplicate_selection=True,
    header_row=None,
    renamed_fio_file=None,
    include_unresolved_duplicates=False,
    sheet_name=None,
    log=None,
    parent=None,
) -> tuple[list[dict], dict]:
    """Проверка файла.

    Returns:
        (подсвеченные строки для таблицы UI, сводка stats)
    """

    def write(msg=""):
        if log:
            log.insert("end", msg + "\n")
            log.see("end")
        print(msg)

    validate_output_not_input(
        output_file,
        sprav_file=sprav_file,
        check_file=check_file,
        renamed_fio_file=renamed_fio_file,
    )

    write("Чтение справочника...")
    spr_df = read_sprav_file(sprav_file)
    schema_warnings = validate_sprav_schema(spr_df)
    for col in schema_warnings:
        write(f"Предупреждение: в справочнике нет колонки «{col}»")

    fio_index = build_fio_index(spr_df)
    write(f"Справочник загружен ({fio_index.person_count()} уникальных ФИО)")

    rename_df = None
    tab_index: dict = {}
    if renamed_fio_file:
        write(f"Загрузка файла смены ФИО: {renamed_fio_file}")
        rename_df = prepare_rename_df(renamed_fio_file)
        if rename_df is None:
            write(
                "В файле смены ФИО нет колонок «Смена ФИО» / «Старое ФИО» / «Новое ФИО» — "
                "автопоиск отключён."
            )
        else:
            tab_index = build_tab_index(spr_df)
            write("Файл смены ФИО загружен")

    # ====== Загрузка файла проверки ======
    wb, ws, values_ws = load_check_file(check_file, sheet_name=sheet_name)
    if header_row is None:
        header_row = find_header_row(ws, values_ws=values_ws)
        write(f"Строка заголовков определена автоматически: {header_row}")
    else:
        header_row = int(header_row)
        write(f"Строка заголовков задана вручную: {header_row}")

    raw_headers = row_values(ws, header_row, values_ws=values_ws)
    headers = [str(v).strip() if v is not None and str(v).strip() != "" else v for v in raw_headers]

    fio_col_index = None
    for idx, h in enumerate(headers, start=1):
        if h is not None and str(h).strip().casefold() == FIO_COL.casefold():
            headers[idx - 1] = FIO_COL
            fio_col_index = idx
            break
    if fio_col_index is None:
        raise ValueError(
            f"В строке {header_row} нет колонки «{FIO_COL}». "
            "Выберите другую строку заголовков."
        )
    data_start = header_row + 1

    # ====== Колонки результата + служебные колонки подсветки дублей ======
    last_col = len(headers)
    extra_cols = [
        COL_HIERARCHY_NAME,
        COL_POSITION_NAME,
        COL_COMMENT_NAME,
        COL_DUP_STEP2_NAME,
        COL_DUP_STEP3_NAME,
        COL_DUP_MANUAL_NAME,
    ]
    ws.insert_cols(last_col + 1, amount=len(extra_cols))
    for offset, name in enumerate(extra_cols, start=1):
        ws.cell(row=header_row, column=last_col + offset, value=name)

    new_col_hierarchy_index = last_col + 1
    new_col_position_index = last_col + 2
    new_col_comment_index = last_col + 3
    dup_step2_col = last_col + 4
    dup_step3_col = last_col + 5
    dup_manual_col = last_col + 6
    headers = headers + extra_cols

    red_fill = get_fill(RED_FILL_COLOR)
    orange_fill = get_fill(ORANGE_FILL_COLOR)
    yellow_fill = get_fill(YELLOW_FILL_COLOR)
    green_fill = get_fill(GREEN_FILL_COLOR)

    write("Проверка сотрудников...")

    red_not_found = []
    orange_dismissed = []
    yellow_duplicates = {}
    include_in_report = []
    repit_not_found = set()

    include_all_missing = False
    exclude_all_missing = False
    output_path = Path(output_file)
    report_path = output_path.with_name(f"{output_path.stem}_included.xlsx")
    report_file = str(report_path)

    # ====== Первый проход ======
    max_data_row = ws.max_row or data_start
    fio_by_row = {
        row: value
        for row, value in iter_column_values(
            ws, fio_col_index, data_start, max_data_row, values_ws
        )
    }
    next_col_by_row = {}
    if fio_col_index + 1 <= (ws.max_column or 1):
        next_col_by_row = {
            row: value
            for row, value in iter_column_values(
                ws, fio_col_index + 1, data_start, max_data_row, values_ws
            )
        }

    for row in range(data_start, max_data_row + 1):
        fio_value = fio_by_row.get(row)
        fio_cell = ws.cell(row=row, column=fio_col_index)
        if isinstance(fio_cell.value, str) and fio_cell.value.startswith("=") and fio_value is not None:
            fio_cell.value = fio_value
        fio_norm = normalize_fio(fio_value) or full_fio_key(fio_value)
        pos_or_hierarchy_value = next_col_by_row.get(row)

        if not fio_value or str(fio_value).strip() == "":
            continue

        records = lookup_fio(fio_index, fio_value)
        via_rename = False

        if not records and rename_df is not None:
            resolved, _via = resolve_via_rename(
                str(fio_value), rename_df, fio_index, tab_index
            )
            if resolved:
                records = resolved
                via_rename = True
                write(f"  {fio_value} — найден через файл смены ФИО → {records[0].get('ФИО', '')}")

        if not records:
            fio_cell.fill = red_fill
            ws.cell(row=row, column=new_col_hierarchy_index, value="")
            ws.cell(row=row, column=new_col_position_index, value="")
            _write_comment(ws, row, new_col_comment_index, MSG_NOT_FOUND)

            if fio_norm not in repit_not_found:
                red_not_found.append((row, fio_value))

                if parent:
                    if exclude_all_missing:
                        write(f"  {fio_value} — не включён в отчёт")
                    elif include_all_missing:
                        include_in_report.append(row)
                    else:
                        answer = _ask_include(parent, fio_value)
                        if answer == "yes_all":
                            include_all_missing = True
                            include_in_report.append(row)
                        elif answer == "no_all":
                            exclude_all_missing = True
                            write(f"  {fio_value} — не включён в отчёт")
                        elif answer == "yes":
                            include_in_report.append(row)
                        else:
                            write(f"  {fio_value} — не включён в отчёт")

                repit_not_found.add(fio_norm)
            elif include_all_missing and not exclude_all_missing:
                include_in_report.append(row)
            continue

        effective_norm = (
            person_group_key(records[0].get("ФИО", fio_value)) or fio_norm
        )

        if len(records) > 1:
            if not check_duplicates:
                chosen = pick_default_duplicate_record(records)
                if chosen is None:
                    # Нет активной основной записи — не пускаем в итог
                    non_joint = [r for r in records if not is_joint(r)]
                    fio_cell.fill = orange_fill
                    if not non_joint:
                        comment = MSG_DUP_JOINTS_ONLY
                    else:
                        comment = MSG_DUP_ALL_DISMISSED
                    _write_comment(ws, row, new_col_comment_index, comment)
                    _write_sprav_columns(
                        ws,
                        row,
                        records[0],
                        new_col_hierarchy_index,
                        new_col_position_index,
                    )
                    orange_dismissed.append(str(fio_value))
                    continue

                fio_cell.fill = yellow_fill
                _write_comment(ws, row, new_col_comment_index, MSG_DUP_DEFAULT)
                _write_sprav_columns(
                    ws, row, chosen, new_col_hierarchy_index, new_col_position_index
                )
                continue

            fio_cell.fill = yellow_fill
            yellow_duplicates[(row, effective_norm, fio_value, pos_or_hierarchy_value)] = records
            _write_comment(ws, row, new_col_comment_index, MSG_DUP_STEP1)
            continue

        if check_dismissed and is_dismissed(records[0]):
            fio_cell.fill = orange_fill
            orange_dismissed.append(str(fio_value))
            _write_comment(ws, row, new_col_comment_index, MSG_DISMISSED)
        elif via_rename:
            _write_comment(ws, row, new_col_comment_index, MSG_RENAMED_FOUND)

        _write_sprav_columns(
            ws, row, records[0], new_col_hierarchy_index, new_col_position_index
        )

    # ====== Новый XLSX для выбранных красных ======
    if include_in_report:
        new_wb = openpyxl.Workbook()
        new_ws = new_wb.active
        for col_idx, header in enumerate(headers, start=1):
            new_ws.cell(row=1, column=col_idx, value=header)
        for i, row_idx in enumerate(include_in_report, start=2):
            for col_idx, header in enumerate(headers, start=1):
                new_ws.cell(row=i, column=col_idx, value=ws.cell(row=row_idx, column=col_idx).value)
        try:
            save_file(new_wb, report_path)
        except (PermissionError, OSError) as exc:
            write(f"Не удалось сохранить файл выбранных сотрудников: {exc}")
        else:
            write(f"Сгенерирован файл с выбранными сотрудниками: {report_file}")
    elif report_path.exists():
        try:
            report_path.unlink()
            write(f"Файл отчёта удалён — ни один сотрудник не включён: {report_file}")
        except PermissionError:
            write(
                f"Не удалось удалить старый файл отчёта (закройте его в Excel): {report_file}"
            )

    # ====== Лог: красные + доп. диагностика по уже загруженному rename_df ======
    if red_not_found:
        write("\nСотрудники, отсутствующие в справочнике (красным):")
        for _, fio in red_not_found:
            write(f"  - {fio}")

        if rename_df is not None:
            write("\nДополнительная проверка оставшихся по файлу смены ФИО...")
            missing_fios = [fio for _, fio in red_not_found]
            change_fio(
                missing_fios,
                write=write,
                fio_index=fio_index,
                spr_df=spr_df,
                rename_df=rename_df,
                tab_index=tab_index,
            )

    # ====== Дубликаты ======
    if check_duplicates and yellow_duplicates:
        write("\nРазбор дубликатов в справочнике...")
        new_d, not_d, dismissed_dups, joints_only = check_duplicate(yellow_duplicates)

        paint_rows_by_dict(ws, new_d, yellow_fill, dup_step2_col)
        paint_rows_by_dict(ws, not_d, green_fill, dup_step2_col)
        _comment_rows(ws, new_d, new_col_comment_index, MSG_DUP_STEP2_YELLOW)
        _comment_rows(ws, not_d, new_col_comment_index, MSG_DUP_STEP2_GREEN)
        write_position_hierarchy(ws, not_d, new_col_hierarchy_index, new_col_position_index)

        paint_rows_by_dict(ws, dismissed_dups, orange_fill, fio_col_index)
        _comment_rows(ws, dismissed_dups, new_col_comment_index, MSG_DUP_ALL_DISMISSED)
        for (row_idx, _, _, _), records in dismissed_dups.items():
            _write_sprav_columns(
                ws, row_idx, records[0], new_col_hierarchy_index, new_col_position_index
            )

        paint_rows_by_dict(ws, joints_only, orange_fill, fio_col_index)
        _comment_rows(ws, joints_only, new_col_comment_index, MSG_DUP_JOINTS_ONLY)
        for (row_idx, _, _, _), records in joints_only.items():
            _write_sprav_columns(
                ws, row_idx, records[0], new_col_hierarchy_index, new_col_position_index
            )

        resolved_count = len(not_d)
        remaining_d, resolved_full = second_check_duplicate(new_d)
        paint_rows_by_dict(ws, remaining_d, yellow_fill, dup_step3_col)
        paint_rows_by_dict(ws, resolved_full, green_fill, dup_step3_col)
        _comment_rows(ws, remaining_d, new_col_comment_index, MSG_DUP_STEP3_YELLOW)
        _comment_rows(ws, resolved_full, new_col_comment_index, MSG_DUP_STEP3_GREEN)
        write_position_hierarchy(ws, resolved_full, new_col_hierarchy_index, new_col_position_index)

        write(
            f"Дубликаты: осталось неоднозначных={len(remaining_d)}, "
            f"однозначно найдено={resolved_count + len(resolved_full)}, "
            f"только уволенные={len(dismissed_dups)}, "
            f"только совместители={len(joints_only)}"
        )

        skip_duplicate_selection = not prompt_duplicate_selection
        seen = set()
        for (row_idx, fio_norm, fio_value, _), records in remaining_d.items():
            if fio_norm in seen:
                continue
            seen.add(fio_norm)
            write("\n" + "=" * 90)
            write(f"Дубликаты в справочнике для: {fio_value}")
            write("=" * 90)
            for idx, rec in enumerate(records, start=1):
                write(f"\nЗапись #{idx}")
                for col, value in rec.items():
                    write(f"  {col}: {value}")

            if parent and not skip_duplicate_selection:
                choice = _choose_duplicate(parent, fio_value, records)
                if choice is SKIP_ALL_DUPLICATES:
                    skip_duplicate_selection = True
                    write("Дальнейший выбор дубликатов пропущен — отчёт сформирован без ручного выбора.")
                elif choice is not None:
                    write(f"Выбрана актуальная запись: {choice}")
                    fill = orange_fill if is_dismissed(choice) else green_fill
                    comment = MSG_DUP_MANUAL_DISMISSED if is_dismissed(choice) else MSG_DUP_MANUAL_ACTIVE
                    for (r, fn, _, _), _recs in list(remaining_d.items()):
                        if fn == fio_norm:
                            ws.cell(row=r, column=dup_manual_col).fill = fill
                            _write_comment(ws, r, new_col_comment_index, comment)
                            _write_sprav_columns(
                                ws,
                                r,
                                choice,
                                new_col_hierarchy_index,
                                new_col_position_index,
                            )
        if remaining_d and skip_duplicate_selection and not prompt_duplicate_selection:
            write("Дубликаты записаны в журнал без ручного выбора.")
    elif check_duplicates:
        write("\nДубликатов не найдено")

    if orange_dismissed:
        write("\nСотрудники, помеченные как уволенные (оранжевым):")
        for fio in sorted(set(orange_dismissed)):
            write(f"  - {fio}")

    if include_unresolved_duplicates:
        for row in range(data_start, ws.max_row + 1):
            comment = ws.cell(row=row, column=new_col_comment_index).value
            text = "" if comment is None else str(comment).strip()
            if text in (MSG_DUP_STEP1, MSG_DUP_STEP2_YELLOW, MSG_DUP_STEP3_YELLOW):
                _write_comment(ws, row, new_col_comment_index, MSG_DUP_UNRESOLVED_INCLUDED)

    results = collect_highlighted_rows(
        ws,
        fio_col=fio_col_index,
        hierarchy_col=new_col_hierarchy_index,
        position_col=new_col_position_index,
        comment_col=new_col_comment_index,
        data_start=data_start,
    )

    final_wb = build_final_workbook(
        ws,
        headers=headers,
        fio_col=fio_col_index,
        comment_col=new_col_comment_index,
        data_start=data_start,
        include_unresolved_duplicates=include_unresolved_duplicates,
    )
    save_file(final_wb, output_file)
    write(f"\nГотово! Итоговый файл (без подсветки): {output_file}")

    review_path = output_path.with_name(f"{output_path.stem}_review.xlsx")
    save_file(wb, str(review_path))
    write(f"Файл с подсветкой для проверки: {review_path}")
    write(f"Подсвеченных строк: {len(results)}")

    stats = summarize_check_rows(
        ws,
        fio_col=fio_col_index,
        comment_col=new_col_comment_index,
        data_start=data_start,
    )
    stats["final_rows"] = max((final_wb.active.max_row or 1) - 1, 0)
    stats["excluded"] = (
        stats["not_found"]
        + stats["dismissed"]
        + (0 if include_unresolved_duplicates else stats["unresolved_dups"])
    )
    write(
        f"Итог: в файле результата {stats['final_rows']} строк; "
        f"не вошло {stats['excluded']} "
        f"(не найдены={stats['not_found']}, уволены={stats['dismissed']}, "
        f"неразрешённые дубли={0 if include_unresolved_duplicates else stats['unresolved_dups']})"
    )
    return results, stats
