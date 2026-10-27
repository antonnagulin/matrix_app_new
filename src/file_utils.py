"""Утилиты чтения/записи Excel с устойчивостью к типичным проблемам."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

from .constants import FIO_COL

# Защита от слишком больших / битых файлов
MAX_XLSX_BYTES = 200 * 1024 * 1024  # 200 МБ
MAX_WORKBOOK_ROWS = 300_000
MAX_WORKBOOK_COLS = 256
MAX_HEADER_SCAN_ROWS = 500


def assert_safe_xlsx_path(file_path: str | Path, *, max_bytes: int = MAX_XLSX_BYTES) -> Path:
    """Проверить размер файла до открытия."""
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"Файл не найден:\n{path}")
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise OSError(f"Не удалось прочитать файл:\n{path}\n{exc}") from exc
    if size > max_bytes:
        mb = size / (1024 * 1024)
        limit_mb = max_bytes / (1024 * 1024)
        raise ValueError(
            f"Файл слишком большой ({mb:.1f} МБ, лимит {limit_mb:.0f} МБ):\n{path}\n"
            "Разбейте данные или увеличьте лимит осознанно."
        )
    if size == 0:
        raise ValueError(f"Файл пустой:\n{path}")
    return path


def assert_worksheet_bounds(ws, *, max_rows: int = MAX_WORKBOOK_ROWS, max_cols: int = MAX_WORKBOOK_COLS) -> None:
    """Ограничить размеры листа после открытия."""
    rows = ws.max_row or 0
    cols = ws.max_column or 0
    if rows > max_rows:
        raise ValueError(
            f"Слишком много строк на листе ({rows}, лимит {max_rows})."
        )
    if cols > max_cols:
        raise ValueError(
            f"Слишком много колонок на листе ({cols}, лимит {max_cols})."
        )


def _open_workbook(file_path, *, data_only: bool, read_only: bool = False):
    try:
        return load_workbook(file_path, data_only=data_only, read_only=read_only)
    except PermissionError as exc:
        raise PermissionError(
            f"Не удалось открыть файл (возможно, он открыт в Excel):\n{file_path}"
        ) from exc
    except InvalidFileException as exc:
        raise ValueError(
            f"Файл повреждён или это не Excel (.xlsx):\n{file_path}"
        ) from exc
    except Exception as exc:
        # zipfile.BadZipFile и прочие повреждения
        name = type(exc).__name__
        if name in {"BadZipFile", "ZIPError", "OSError"} or "zip" in str(exc).lower():
            raise ValueError(
                f"Файл повреждён или это не Excel (.xlsx):\n{file_path}\n{exc}"
            ) from exc
        raise


def read_sprav_file(file_path: str, sheet_name: str | None = None) -> pd.DataFrame:
    """Чтение справочника в DataFrame."""
    assert_safe_xlsx_path(file_path)
    try:
        df = pd.read_excel(file_path, sheet_name=sheet_name or 0, dtype=str).fillna("")
    except PermissionError as exc:
        raise PermissionError(
            f"Не удалось прочитать файл (возможно, он открыт в Excel):\n{file_path}"
        ) from exc
    except Exception as exc:
        raise ValueError(
            f"Не удалось прочитать справочник (повреждённый или неверный формат):\n"
            f"{file_path}\n{exc}"
        ) from exc
    if len(df) > MAX_WORKBOOK_ROWS:
        raise ValueError(
            f"В справочнике слишком много строк ({len(df)}, лимит {MAX_WORKBOOK_ROWS})."
        )
    return df


def list_sheet_names(file_path: str) -> list[str]:
    """Имена листов книги."""
    assert_safe_xlsx_path(file_path)
    wb = _open_workbook(file_path, data_only=False, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _pick_sheet(wb, sheet_name: str | None):
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            available = ", ".join(wb.sheetnames) if wb.sheetnames else "(пусто)"
            raise ValueError(
                f"Лист «{sheet_name}» не найден в файле.\n"
                f"Доступные листы: {available}"
            )
        return wb[sheet_name]
    return wb.active


def load_check_file(file_path, sheet_name: str | None = None):
    """
    Открыть файл проверки для правки (подсветка).

    Дополнительно грузится data_only-копия для чтения значений формул (если Excel
    сохранил кэш). Возвращает (wb, ws, values_ws|None).
    """
    assert_safe_xlsx_path(file_path)
    wb = _open_workbook(file_path, data_only=False)
    ws = _pick_sheet(wb, sheet_name)
    assert_worksheet_bounds(ws)

    values_ws = None
    try:
        wb_values = _open_workbook(file_path, data_only=True)
        values_ws = _pick_sheet(wb_values, sheet_name)
        # Держим wb_values живым через атрибут, иначе GC закроет
        wb._matrix_values_wb = wb_values  # noqa: SLF001
    except ValueError:
        raise
    except Exception:
        values_ws = None

    return wb, ws, values_ws


def save_file(wb, file_path) -> None:
    """Сохранить файл Excel с понятной ошибкой при блокировке."""
    try:
        wb.save(file_path)
    except PermissionError as exc:
        raise PermissionError(
            f"Не удалось сохранить файл (закройте его в Excel, если открыт):\n{file_path}"
        ) from exc
    except OSError as exc:
        raise OSError(f"Ошибка записи файла:\n{file_path}\n{exc}") from exc


def get_fill(color_hex: str) -> PatternFill:
    """Получить цвет заливки ячейки."""
    return PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")


def clear_fill() -> PatternFill:
    """Сбросить заливку ячейки."""
    return PatternFill(fill_type=None)


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _merged_origin_value(ws, row: int, col: int):
    """Если ячейка в объединении — значение из левого верхнего угла."""
    cell = ws.cell(row=row, column=col)
    if cell.value is not None:
        return cell.value
    # merged_cells может быть недоступен в read_only
    merged = getattr(ws, "merged_cells", None)
    if not merged or not getattr(merged, "ranges", None):
        return None
    coord = cell.coordinate
    for rng in merged.ranges:
        if coord in rng:
            min_col, min_row, _max_col, _max_row = range_boundaries(str(rng))
            return ws.cell(row=min_row, column=min_col).value
    return None


def cell_value(ws, row: int, col: int, values_ws=None):
    """
    Прочитанное значение ячейки с учётом merge и кэша формул (data_only).
    """
    raw = _merged_origin_value(ws, row, col)
    if isinstance(raw, str) and raw.startswith("=") and values_ws is not None:
        cached = _merged_origin_value(values_ws, row, col)
        if cached is not None:
            return cached
    if raw is None and values_ws is not None:
        cached = _merged_origin_value(values_ws, row, col)
        if cached is not None:
            return cached
    return raw


def row_values(ws, row: int, max_col: int | None = None, values_ws=None) -> list:
    """Значения ячеек строки (1-based)."""
    last_col = max_col or ws.max_column or 1
    return [cell_value(ws, row, c, values_ws) for c in range(1, last_col + 1)]


def iter_column_values(ws, col: int, start_row: int, end_row: int, values_ws=None):
    """
    Итератор (row, value) по колонке. Без merge — через iter_rows (быстрее).
    """
    has_merges = bool(getattr(getattr(ws, "merged_cells", None), "ranges", None))
    if not has_merges:
        for row_idx, (raw,) in enumerate(
            ws.iter_rows(
                min_row=start_row,
                max_row=end_row,
                min_col=col,
                max_col=col,
                values_only=True,
            ),
            start=start_row,
        ):
            if isinstance(raw, str) and raw.startswith("=") and values_ws is not None:
                raw = values_ws.cell(row=row_idx, column=col).value
            elif raw is None and values_ws is not None:
                cached = values_ws.cell(row=row_idx, column=col).value
                if cached is not None:
                    raw = cached
            yield row_idx, raw
        return

    for row in range(start_row, end_row + 1):
        yield row, cell_value(ws, row, col, values_ws)


def find_header_row(
    ws,
    marker: str = FIO_COL,
    *,
    max_scan: int | None = None,
    values_ws=None,
) -> int:
    """
    Найти номер строки (1-based) с заголовками по маркеру (обычно «ФИО»).
    """
    limit = max_scan if max_scan is not None else min(
        max(ws.max_row or 1, 1), MAX_HEADER_SCAN_ROWS
    )
    marker_norm = marker.strip().casefold()
    max_col = min(ws.max_column or 1, MAX_WORKBOOK_COLS)
    has_merges = bool(getattr(getattr(ws, "merged_cells", None), "ranges", None))

    if not has_merges:
        for row_idx, row_vals in enumerate(
            ws.iter_rows(
                min_row=1,
                max_row=limit,
                min_col=1,
                max_col=max_col,
                values_only=True,
            ),
            start=1,
        ):
            for col_idx, raw in enumerate(row_vals, start=1):
                if isinstance(raw, str) and raw.startswith("=") and values_ws is not None:
                    raw = values_ws.cell(row=row_idx, column=col_idx).value
                text = _cell_text(raw)
                if text.casefold() == marker_norm:
                    return row_idx
    else:
        for row in range(1, limit + 1):
            for col in range(1, max_col + 1):
                text = _cell_text(cell_value(ws, row, col, values_ws))
                if text.casefold() == marker_norm:
                    return row

    raise ValueError(
        f"Не найдена строка заголовков с колонкой «{marker}» "
        f"(проверены строки 1–{limit})."
    )


def peek_sheet_rows(
    file_path: str,
    max_rows: int = 40,
    sheet_name: str | None = None,
) -> tuple[list[list], int]:
    """
    Прочитать первые строки файла проверки для превью в UI.

    Returns:
        (rows, detected_header_row)
    """
    wb, ws, values_ws = load_check_file(file_path, sheet_name=sheet_name)
    try:
        detected = find_header_row(ws, values_ws=values_ws)
    except ValueError:
        detected = 0

    last = min(max_rows, ws.max_row or 1)
    max_col = min(ws.max_column or 1, 12)
    rows = []
    for r in range(1, last + 1):
        values = [_cell_text(v) for v in row_values(ws, r, max_col, values_ws)]
        rows.append([r, *values])
    return rows, detected
