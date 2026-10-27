"""Сборка итогового workbook и сводки по строкам проверки."""

from __future__ import annotations

import openpyxl

COL_HIERARCHY_NAME = "Подразделение"
COL_POSITION_NAME = "Штатная должность (справочник)"
COL_COMMENT_NAME = "Комментарий"
COL_DUP_STEP2_NAME = "Дубль шаг 2"
COL_DUP_STEP3_NAME = "Дубль шаг 3"
COL_DUP_MANUAL_NAME = "Дубль выбор"

FINAL_EXCLUDE_COLUMNS = {
    COL_COMMENT_NAME,
    COL_DUP_STEP2_NAME,
    COL_DUP_STEP3_NAME,
    COL_DUP_MANUAL_NAME,
}

MSG_NOT_FOUND = "Не найден в справочнике"
MSG_DISMISSED = "Сотрудник уволен или неактивен"
MSG_DUP_STEP1 = "Дубликат в справочнике (шаг 1)"
MSG_DUP_STEP1_NO_CHECK = "Дубликат в справочнике (разбор дубликатов отключён)"
MSG_DUP_DEFAULT = "Дубликат не разбирался — взята запись по умолчанию"
MSG_DUP_STEP2_YELLOW = "Дубликат не разрешён после 1-й проверки (шаг 2)"
MSG_DUP_STEP2_GREEN = "Дубликат разрешён после 1-й проверки (шаг 2)"
MSG_DUP_STEP3_YELLOW = "Дубликат не разрешён после 2-й проверки (шаг 3)"
MSG_DUP_STEP3_GREEN = "Дубликат разрешён по полному ФИО (шаг 3)"
MSG_DUP_ALL_DISMISSED = "Все записи дубликата уволены"
MSG_DUP_JOINTS_ONLY = "Только внутренние совместители — нет основной записи"
MSG_DUP_MANUAL_ACTIVE = "Дубликат выбран вручную (шаг 3)"
MSG_DUP_MANUAL_DISMISSED = "Дубликат выбран вручную — уволенная запись (шаг 3)"
MSG_RENAMED_FOUND = "Найден через файл смены ФИО"
MSG_DUP_UNRESOLVED_INCLUDED = "Неразрешённый дубликат (включён в итог по настройке)"

STATUS_RED = "red"
STATUS_ORANGE = "orange"
STATUS_YELLOW = "yellow"
STATUS_GREEN = "green"

COMMENT_STATUS = {
    MSG_NOT_FOUND: STATUS_RED,
    MSG_DISMISSED: STATUS_ORANGE,
    MSG_DUP_STEP1: STATUS_YELLOW,
    MSG_DUP_STEP1_NO_CHECK: STATUS_YELLOW,
    MSG_DUP_DEFAULT: STATUS_GREEN,
    MSG_DUP_STEP2_YELLOW: STATUS_YELLOW,
    MSG_DUP_STEP2_GREEN: STATUS_GREEN,
    MSG_DUP_STEP3_YELLOW: STATUS_YELLOW,
    MSG_DUP_STEP3_GREEN: STATUS_GREEN,
    MSG_DUP_ALL_DISMISSED: STATUS_ORANGE,
    MSG_DUP_JOINTS_ONLY: STATUS_ORANGE,
    MSG_DUP_MANUAL_ACTIVE: STATUS_GREEN,
    MSG_DUP_MANUAL_DISMISSED: STATUS_ORANGE,
    MSG_RENAMED_FOUND: STATUS_GREEN,
    MSG_DUP_UNRESOLVED_INCLUDED: STATUS_YELLOW,
}

UNRESOLVED_DUP_COMMENTS = {
    MSG_DUP_STEP1,
    MSG_DUP_STEP2_YELLOW,
    MSG_DUP_STEP3_YELLOW,
    MSG_DUP_UNRESOLVED_INCLUDED,
}


def status_from_comment(comment) -> str | None:
    """Цвет строки таблицы по тексту комментария; None — без подсветки."""
    if comment is None:
        return None
    text = str(comment).strip()
    if not text:
        return None
    return COMMENT_STATUS.get(text)


def is_accepted_for_final(comment, *, include_unresolved_duplicates: bool = False) -> bool:
    """Строка попадает в итоговый файл: активный / разрешённый дубликат / опционально неразрешённый."""
    if comment is None or str(comment).strip() == "":
        return True
    text = str(comment).strip()
    status = COMMENT_STATUS.get(text)
    if status == STATUS_GREEN:
        return True
    if include_unresolved_duplicates and text in UNRESOLVED_DUP_COMMENTS:
        return True
    return False


def collect_highlighted_rows(
    ws,
    *,
    fio_col: int,
    hierarchy_col: int,
    position_col: int,
    comment_col: int,
    data_start: int,
) -> list[dict]:
    """Собрать подсвеченные строки результата для таблицы UI."""
    rows: list[dict] = []
    for row in range(data_start, ws.max_row + 1):
        fio_value = ws.cell(row=row, column=fio_col).value
        if not fio_value or str(fio_value).strip() == "":
            continue
        comment = ws.cell(row=row, column=comment_col).value
        status = status_from_comment(comment)
        if status is None:
            continue
        hierarchy = ws.cell(row=row, column=hierarchy_col).value
        position = ws.cell(row=row, column=position_col).value
        rows.append(
            {
                "row": row,
                "fio": str(fio_value).strip(),
                "hierarchy": "" if hierarchy is None else str(hierarchy).strip(),
                "position": "" if position is None else str(position).strip(),
                "comment": "" if comment is None else str(comment).strip(),
                "status": status,
            }
        )
    return rows


def build_final_workbook(
    ws,
    *,
    headers: list,
    fio_col: int,
    comment_col: int,
    data_start: int,
    include_unresolved_duplicates: bool = False,
):
    """
    Итоговый файл без подсветки: активные и разрешённые дубликаты.
    Служебные колонки комментария/дублей не копируются.
    """
    keep_indexes = [
        i for i, h in enumerate(headers, start=1) if h not in FINAL_EXCLUDE_COLUMNS
    ]
    final_headers = [headers[i - 1] for i in keep_indexes]
    final_wb = openpyxl.Workbook()
    final_ws = final_wb.active
    for col_idx, header in enumerate(final_headers, start=1):
        final_ws.cell(row=1, column=col_idx, value=header)

    out_row = 2
    for row in range(data_start, ws.max_row + 1):
        fio_value = ws.cell(row=row, column=fio_col).value
        if not fio_value or str(fio_value).strip() == "":
            continue
        comment = ws.cell(row=row, column=comment_col).value
        if not is_accepted_for_final(
            comment, include_unresolved_duplicates=include_unresolved_duplicates
        ):
            continue

        for out_col, src_col in enumerate(keep_indexes, start=1):
            final_ws.cell(
                row=out_row,
                column=out_col,
                value=ws.cell(row=row, column=src_col).value,
            )
        out_row += 1

    return final_wb


def summarize_check_rows(ws, *, fio_col: int, comment_col: int, data_start: int) -> dict:
    """Сводка по строкам проверки для диалога итога."""
    stats = {
        "total": 0,
        "ok": 0,
        "not_found": 0,
        "dismissed": 0,
        "unresolved_dups": 0,
        "resolved_dups": 0,
        "renamed": 0,
    }
    for row in range(data_start, ws.max_row + 1):
        fio_value = ws.cell(row=row, column=fio_col).value
        if not fio_value or str(fio_value).strip() == "":
            continue
        stats["total"] += 1
        comment = ws.cell(row=row, column=comment_col).value
        text = "" if comment is None else str(comment).strip()
        status = status_from_comment(text)
        if not text:
            stats["ok"] += 1
        elif text == MSG_NOT_FOUND:
            stats["not_found"] += 1
        elif text in (
            MSG_DISMISSED,
            MSG_DUP_ALL_DISMISSED,
            MSG_DUP_JOINTS_ONLY,
            MSG_DUP_MANUAL_DISMISSED,
        ):
            stats["dismissed"] += 1
        elif text in UNRESOLVED_DUP_COMMENTS or text == MSG_DUP_STEP1_NO_CHECK:
            stats["unresolved_dups"] += 1
        elif status == STATUS_GREEN:
            if text == MSG_RENAMED_FOUND:
                stats["renamed"] += 1
            else:
                stats["resolved_dups"] += 1
        else:
            stats["ok"] += 1
    return stats
