from ..src.filters import normalize_fio, is_dismissed
from ..src.file_utils import read_sprav_file, load_check_file, save_file, get_fill
from ..src.constants import FIO_COL, RED_FILL_COLOR, ORANGE_FILL_COLOR
import openpyxl
import pandas as pdd
from ..src.check_chainge_fio import chainge_fio


def core(
        fio_col_index,
        ws,
        spr_df,
        red_fill,
        new_col_hierarchy_index,
        new_col_position_index,
        join_with_slash,
        check_dismissed,
        orange_fill,
        yellow_fill,
        write,
        check_duplicates,
        wb,
        output_file,
        chainge_fio_file,
    ):

        # ФИО -> список записей
    fio_dict = {}
    for _, row in spr_df.iterrows():
        fio_norm = normalize_fio(row[FIO_COL])
        if fio_norm:
            fio_dict.setdefault(fio_norm, []).append(row)
    
    
    write(f"Справочник загружен ({len(fio_dict)} уникальных ФИО)")

    red_not_found = []
    orange_dismissed = []

    for row in range(2, ws.max_row + 1):
        fio_cell = ws.cell(row=row, column=fio_col_index)
        fio_value = fio_cell.value
        if not fio_value or str(fio_value).strip() == "":
            continue

        fio_norm = normalize_fio(fio_value)

        # ❌ Нет в справочнике
        if fio_norm not in fio_dict:
            fio_cell.fill = red_fill
            ws.cell(row=row, column=new_col_hierarchy_index, value="")
            ws.cell(row=row, column=new_col_position_index, value="")
            red_not_found.append(str(fio_value))
            continue

        # 🟢 Берем первую запись из справочника
        record = fio_dict[fio_norm][0]
        hierarchy_str = join_with_slash(record)
        position_str = record.get("Штатная должность", "")

        ws.cell(row=row, column=new_col_hierarchy_index, value=hierarchy_str)
        ws.cell(row=row, column=new_col_position_index, value=position_str)
        

        if len(fio_dict[fio_norm]) > 1:
            # print(fio_dict[fio_norm])
            fio_cell_1 = ws.cell(row=row, column=fio_col_index)
            fio_cell_1.fill = yellow_fill

            tmp = list()
            for item in fio_dict[fio_norm]:
                tmp.append(item['Таб.№'])
            

            dub = ws.cell(row=row, column=new_col_position_index+1, value=str(tmp))
            dub.fill = yellow_fill

            tmp=list()

        # 🟠 Проверка увольнения
        if check_dismissed:
            records = fio_dict[fio_norm]
            if len(records) == 1:
                if any(is_dismissed(rec) for rec in records):
                    fio_cell.fill = orange_fill
                    orange_dismissed.append(str(fio_value))


    chainge_fio(red_not_found, chainge_fio_file)

    # print("Не нашлись в справочнике")
    # print(red_not_found)

    # chainge_fio_df = read_sprav_file(chainge_fio_file)

    # # Предполагаем, что ФИО лежит в колонке "ФИО"
    # # Нормализуем ФИО в справочнике ОДИН раз
    # chainge_fio_df["ФИО_NORM"] = chainge_fio_df["Смена ФИО"].apply(normalize_fio)
    # not_in_all = []
    # for item in red_not_found:
    #     item_norm = normalize_fio(item)

    #     matched_rows = chainge_fio_df[
    #         chainge_fio_df["ФИО_NORM"] == item_norm
    #     ]

    #     if not matched_rows.empty:
    #         print(f"\nНайдено совпадение для: {item}")
    #         print(matched_rows.drop(columns=["ФИО_NORM"]))

    #     else:
    #         not_in_all.append(item)

    # if not_in_all:        
    #     print("Не нашлись вообще")        
    #     print(not_in_all)

        
    # ====== Лог: отсутствующие ======
    # write()
    # if red_not_found:
    #     write("Сотрудники, отсутствующие в справочнике (красным):")
    #     for fio in sorted(set(red_not_found)):
    #         write(f"  - {fio}")
    # else:
    #     write("Сотрудников, отсутствующих в справочнике, не найдено")

    # # ====== Лог: дубликаты ======
    # if check_duplicates:
    #     write()
    #     write("Поиск дубликатов в справочнике (для уволенных)...")
    #     seen = set()
    #     found_any = False

    #     for fio in sorted(set(orange_dismissed)):
    #         fio_norm = normalize_fio(fio)
    #         records = fio_dict.get(fio_norm, [])
    #         if len(records) > 1 and fio_norm not in seen:
    #             seen.add(fio_norm)
    #             found_any = True
    #             write()
    #             write("=" * 90)
    #             write(f"Дубликаты в справочнике для: {fio}")
    #             write("=" * 90)
    #             for idx, rec in enumerate(records, start=1):
    #                 write(f"\nЗапись #{idx}")
    #                 for col, value in rec.items():
    #                     write(f"  {col}: {value}")

    #     if not found_any:
    #         write("Дубликатов не найдено")

    save_file(wb, output_file)
    write()
    write(f"Готово! Файл сохранён: {output_file}")