from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from collections import Counter

# === НАСТРОЙКИ ===
file_path = "/home/wol/programs/scripts/matrix_applications/данные/результат1.xlsx"      # путь к файлу
column_letter = "D"           # столбец для проверки
sheet_name = "Лист1"         # имя листа


seen = set()  # сюда будем записывать уже встреченные значения

# === Загрузка файла ===
wb = load_workbook(file_path)
ws = wb[sheet_name]

purple_fill = PatternFill(
    start_color="800080",
    end_color="800080",
    fill_type="solid"
)

# идём сверху вниз
for row in range(1, ws.max_row + 1):
    cell_value = ws[f"{column_letter}{row}"].value

    if cell_value in seen:
        # если уже встречалось — красим строку
        for col in range(1, ws.max_column + 1):
            ws.cell(row=row, column=col).fill = purple_fill
    else:
        # если видим первый раз — запоминаем
        seen.add(cell_value)

# === Сохранение ===
wb.save("output.xlsx")

print("Готово! Дубликаты выделены фиолетовым.")
