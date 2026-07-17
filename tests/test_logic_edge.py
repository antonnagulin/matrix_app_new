"""Дополнительные тесты run_check — граничные случаи."""

import pytest
import pandas as pd
from openpyxl import Workbook

from src.logic import run_check, SKIP_ALL_DUPLICATES
from src.constants import RED_FILL_COLOR, ORANGE_FILL_COLOR, ELOYY_FILL_COLOR, GREEN_FILL_COLOR


class DummyLog:
    def __init__(self):
        self.lines = []

    def insert(self, *_args):
        self.lines.append(_args[1])

    def see(self, *_):
        pass

    def text(self):
        return "".join(self.lines)


def _patch_io(monkeypatch, sprav_df, wb, ws):
    monkeypatch.setattr("src.logic.read_sprav_file", lambda _: sprav_df)
    monkeypatch.setattr("src.logic.load_check_file", lambda _: (wb, ws))
    monkeypatch.setattr("src.logic.save_file", lambda *_: None)


def test_invalid_header_row_raises(monkeypatch):
    sprav_df = pd.DataFrame([{"ФИО": "Иванов Иван Иванович", "Увольн.": "31.12.9999"}])
    wb = Workbook()
    ws = wb.active
    ws["A5"] = "Не ФИО"
    ws["A6"] = "Иванов"
    log = DummyLog()
    _patch_io(monkeypatch, sprav_df, wb, ws)

    with pytest.raises(ValueError, match="нет колонки"):
        run_check(
            "sprav.xlsx",
            "check.xlsx",
            "out.xlsx",
            header_row=5,
            log=log,
        )


def test_duplicate_yellow_without_duplicate_check(monkeypatch):
    """При len>1 в справочнике — жёлтый даже если check_duplicates=False."""
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Менеджер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
        ]
    )
    wb = Workbook()
    ws = wb.active
    ws.append(["ФИО", "Отдел"])
    ws.append(["Иванов И. И.", "IT"])
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=False,
        check_duplicates=False,
    )

    assert ELOYY_FILL_COLOR in (ws["A2"].fill.start_color.rgb or "")


def test_duplicate_choice_marks_green(monkeypatch):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Смирнов Алексей Николаевич",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
            {
                "ФИО": "Смирнов Алексей Николаевич",
                "Штатная должность": "Программист",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
        ]
    )
    wb = Workbook()
    ws = wb.active
    ws.append(["ФИО", "Отдел"])
    ws.append(["Смирнов Алексей Николаевич", "IT"])
    active = sprav_df.iloc[0].to_dict()

    monkeypatch.setattr("src.logic.choose_actual_record", lambda *_a, **_k: active)
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=True,
        check_duplicates=True,
        parent=object(),
    )

    assert GREEN_FILL_COLOR in (ws["A2"].fill.start_color.rgb or "")


def test_duplicate_choice_marks_orange_when_dismissed(monkeypatch):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Смирнов Алексей Николаевич",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
            {
                "ФИО": "Смирнов Алексей Николаевич",
                "Штатная должность": "Программист",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
        ]
    )
    wb = Workbook()
    ws = wb.active
    ws.append(["ФИО", "Отдел"])
    ws.append(["Смирнов Алексей Николаевич", "IT"])
    dismissed = {
        "ФИО": "Смирнов Алексей Николаевич",
        "Штатная должность": "Интеграция: штатная должность по умолч.",
        "Статус занятости": "уволился",
        "Увольн.": "01.01.2020",
    }

    monkeypatch.setattr("src.logic.choose_actual_record", lambda *_a, **_k: dismissed)
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=True,
        check_duplicates=True,
        parent=object(),
    )

    assert ORANGE_FILL_COLOR in (ws["A2"].fill.start_color.rgb or "")


def test_skip_all_constant_value():
    assert SKIP_ALL_DUPLICATES == "skip_all"
