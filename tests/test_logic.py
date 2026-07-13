"""Интеграционный тест run_check на реальных xlsx-файлах."""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.logic import run_check
from src.constants import RED_FILL_COLOR, ORANGE_FILL_COLOR


def test_run_check_end_to_end(tmp_path):
    spr_file = tmp_path / "sprav.xlsx"
    pd.DataFrame(
        {
            "ФИО": ["Иванов Иван Иванович", "Петров Пётр Петрович"],
            "Штатная должность": [
                "Менеджер",
                "Интеграция: штатная должность по умолч.",
            ],
            "Статус занятости": ["работает", "работает"],
            "Увольн.": ["31.12.9999", "31.12.9999"],
        }
    ).to_excel(spr_file, index=False)

    check_file = tmp_path / "check.xlsx"
    pd.DataFrame(
        {
            "ФИО": ["Иванов И. И.", "Петров П. П.", "Сидоров С. С."],
            "Должность": ["Менеджер", "Менеджер", "Аналитик"],
        }
    ).to_excel(check_file, index=False)

    output_file = tmp_path / "output.xlsx"

    run_check(
        sprav_file=spr_file,
        check_file=check_file,
        output_file=output_file,
        check_dismissed=True,
        check_duplicates=False,
        log=None,
        parent=None,
    )

    assert Path(output_file).exists()
    ws = load_workbook(output_file).active

    # Иванов — активен, без подсветки
    ivanov_fill = ws["A2"].fill.start_color.rgb or ""
    assert RED_FILL_COLOR not in ivanov_fill
    assert ORANGE_FILL_COLOR not in ivanov_fill

    # Петров — уволен по должности → оранжевый
    assert ORANGE_FILL_COLOR in (ws["A3"].fill.start_color.rgb or "")

    # Сидоров — нет в справочнике → красный
    assert RED_FILL_COLOR in (ws["A4"].fill.start_color.rgb or "")
