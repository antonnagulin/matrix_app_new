"""Интеграционный тест на sample_data/check_medium.xlsx."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.logic import run_check
from src.constants import RED_FILL_COLOR, ELOYY_FILL_COLOR, GREEN_FILL_COLOR

ROOT = Path(__file__).resolve().parents[1]
SPRAV = ROOT / "sample_data" / "sprav_medium.xlsx"
CHECK = ROOT / "sample_data" / "check_medium.xlsx"


@pytest.mark.skipif(not SPRAV.is_file() or not CHECK.is_file(), reason="medium sample missing")
def test_medium_sample_produces_expected_highlights(tmp_path):
    output = tmp_path / "medium_result.xlsx"
    log_lines = []

    class Log:
        def insert(self, _idx, text):
            log_lines.append(text)

        def see(self, _idx):
            pass

    run_check(
        str(SPRAV),
        str(CHECK),
        str(output),
        check_dismissed=True,
        check_duplicates=True,
        prompt_duplicate_selection=False,
        log=Log(),
    )

    assert output.exists()
    review = tmp_path / "medium_result_review.xlsx"
    assert review.exists()
    ws = load_workbook(review).active

    def _has_color(row: int, col: int, color: str) -> bool:
        return color in (ws.cell(row=row, column=col).fill.start_color.rgb or "")

    def _col(header: str) -> int:
        for col in range(1, (ws.max_column or 1) + 1):
            if ws.cell(row=1, column=col).value == header:
                return col
        raise AssertionError(f"Колонка «{header}» не найдена")

    step2 = _col("Дубль шаг 2")
    step3 = _col("Дубль шаг 3")

    # Шаг 1: жёлтый на ФИО; шаг 2: зелёный в служебной колонке (не «Отдел»)
    assert _has_color(17, 1, ELOYY_FILL_COLOR)
    assert _has_color(17, step2, GREEN_FILL_COLOR)
    assert not _has_color(17, 2, GREEN_FILL_COLOR)
    assert _has_color(18, 1, ELOYY_FILL_COLOR)
    assert _has_color(18, step2, GREEN_FILL_COLOR)
    # Два активных — жёлтый на шагах 2 и 3
    assert _has_color(19, 1, ELOYY_FILL_COLOR)
    assert _has_color(19, step2, ELOYY_FILL_COLOR)
    assert _has_color(19, step3, ELOYY_FILL_COLOR)
    assert not _has_color(19, 2, ELOYY_FILL_COLOR)
    assert _has_color(20, 1, ELOYY_FILL_COLOR)
    assert _has_color(20, step2, ELOYY_FILL_COLOR)
    assert _has_color(20, step3, ELOYY_FILL_COLOR)
    assert _has_color(27, 1, RED_FILL_COLOR)

    # Итоговый файл без подсветки и без «проблемных» строк
    clean = load_workbook(output).active
    clean_fios = [
        clean.cell(row=r, column=1).value
        for r in range(2, (clean.max_row or 1) + 1)
        if clean.cell(row=r, column=1).value
    ]
    assert "Смирнов А. Н." not in clean_fios
    assert "Васильева М. И." not in clean_fios
    assert "Неизвестный Н. Н." not in clean_fios
    for r in range(2, (clean.max_row or 1) + 1):
        fill = clean.cell(row=r, column=1).fill.start_color.rgb or ""
        assert RED_FILL_COLOR not in fill
        assert ELOYY_FILL_COLOR not in fill

    log_text = "".join(log_lines)
    assert "Строка заголовков" in log_text
    assert "Разбор дубликатов" in log_text
