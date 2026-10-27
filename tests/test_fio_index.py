"""Тесты многоуровневого индекса ФИО."""

import pandas as pd

from src.fio_index import (
    build_fio_index,
    full_fio_key,
    initials_fio_key,
    lookup_fio,
    pick_default_duplicate_record,
    surname_initial_key,
    surname_name_key,
)
from src.logic import MSG_DUP_DEFAULT, is_accepted_for_final, run_check
from openpyxl import Workbook


def test_keys():
    assert full_fio_key("Иванов Иван Иванович") == "ИВАНОВ ИВАН ИВАНОВИЧ"
    assert initials_fio_key("Иванов Иван Иванович") == "ИВАНОВ ИИ"
    assert initials_fio_key("Иванов И. И.") == "ИВАНОВ ИИ"
    assert surname_name_key("Сидоров Сергей Сергеевич") == "СИДОРОВ СЕРГЕЙ"
    assert surname_name_key("Сидоров Сергей") == "СИДОРОВ СЕРГЕЙ"
    assert surname_initial_key("Иванов Иван Иванович") == "ИВАНОВ И"
    assert surname_initial_key("Иванов И.") == "ИВАНОВ И"
    assert surname_initial_key("Иванов И") == "ИВАНОВ И"


def test_collision_initials_returns_both_people():
    sprav = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
            {
                "ФИО": "Иванов Игорь Иванович",
                "Штатная должность": "Менеджер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
        ]
    )
    index = build_fio_index(sprav)
    found = lookup_fio(index, "Иванов И. И.")
    assert len(found) == 2
    names = {r["ФИО"] for r in found}
    assert names == {"Иванов Иван Иванович", "Иванов Игорь Иванович"}


def test_two_token_initial_with_dot():
    sprav = pd.DataFrame(
        [
            {
                "ФИО": "Козлов Кирилл Кириллович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            }
        ]
    )
    index = build_fio_index(sprav)
    for query in ("Козлов К.", "Козлов К", "Козлов к."):
        found = lookup_fio(index, query)
        assert len(found) == 1, query
        assert found[0]["ФИО"] == "Козлов Кирилл Кириллович"


def test_two_token_initial_ambiguous():
    sprav = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
            {
                "ФИО": "Иванов Игорь Иванович",
                "Штатная должность": "Менеджер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
        ]
    )
    index = build_fio_index(sprav)
    found = lookup_fio(index, "Иванов И.")
    assert len(found) == 2


def test_hyphenated_surname():
    sprav = pd.DataFrame(
        [
            {
                "ФИО": "Петрова-Сидорова Анна Ивановна",
                "Штатная должность": "Бухгалтер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            }
        ]
    )
    index = build_fio_index(sprav)
    for query in (
        "Петрова-Сидорова А. И.",
        "Петрова-Сидорова Анна",
        "Петрова–Сидорова А. И.",  # en-dash
        "Петрова-Сидорова А",
    ):
        found = lookup_fio(index, query)
        assert len(found) == 1, query
        assert found[0]["ФИО"] == "Петрова-Сидорова Анна Ивановна"


def test_full_fio_no_collision():
    sprav = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
            {
                "ФИО": "Иванов Игорь Иванович",
                "Штатная должность": "Менеджер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
        ]
    )
    index = build_fio_index(sprav)
    found = lookup_fio(index, "Иванов Игорь Иванович")
    assert len(found) == 1
    assert found[0]["ФИО"] == "Иванов Игорь Иванович"


def test_two_word_matches_three_word():
    sprav = pd.DataFrame(
        [
            {
                "ФИО": "Сидоров Сергей Сергеевич",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            }
        ]
    )
    index = build_fio_index(sprav)
    found = lookup_fio(index, "Сидоров Сергей")
    assert len(found) == 1
    assert found[0]["ФИО"] == "Сидоров Сергей Сергеевич"


def test_two_word_ambiguous_two_patronymics():
    sprav = pd.DataFrame(
        [
            {
                "ФИО": "Сидоров Сергей Сергеевич",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
            {
                "ФИО": "Сидоров Сергей Петрович",
                "Штатная должность": "Менеджер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
        ]
    )
    index = build_fio_index(sprav)
    found = lookup_fio(index, "Сидоров Сергей")
    assert len(found) == 2


def test_duplicates_off_keeps_row_in_final(monkeypatch):
    sprav = pd.DataFrame(
        [
            {
                "ФИО": "Сидоров Сидор Сидорович",
                "Штатная должность": "Аналитик",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
            {
                "ФИО": "Сидоров Сидор Сидорович",
                "Штатная должность": "Интеграция: штатная должность по умолч.",
                "Статус занятости": "уволился",
                "Увольн.": "01.01.2020",
            },
        ]
    )
    wb = Workbook()
    ws = wb.active
    ws.append(["ФИО", "Отдел"])
    ws.append(["Сидоров С. С.", "Аналитика"])

    saved = {}
    monkeypatch.setattr("src.logic.read_sprav_file", lambda _: sprav)
    monkeypatch.setattr("src.logic.load_check_file", lambda *a, **k: (wb, ws, None))
    monkeypatch.setattr(
        "src.logic.save_file",
        lambda workbook, path: saved.setdefault("files", {}).__setitem__(str(path), workbook),
    )

    results, stats = run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=True,
        check_duplicates=False,
    )
    assert (ws["E2"].value or "") == MSG_DUP_DEFAULT
    assert is_accepted_for_final(MSG_DUP_DEFAULT) is True
    clean = saved["files"]["out.xlsx"].active
    assert clean["A2"].value == "Сидоров С. С."
    assert stats["final_rows"] == 1


def test_duplicates_off_all_dismissed_excluded_from_final(monkeypatch):
    sprav = pd.DataFrame(
        [
            {
                "ФИО": "Сидоров Сидор Сидорович",
                "Штатная должность": "Интеграция: штатная должность по умолч.",
                "Статус занятости": "уволился",
                "Увольн.": "01.01.2020",
            },
            {
                "ФИО": "Сидоров Сидор Сидорович",
                "Штатная должность": "Интеграция: штатная должность по умолч.",
                "Статус занятости": "уволился",
                "Увольн.": "15.03.2021",
            },
        ]
    )
    wb = Workbook()
    ws = wb.active
    ws.append(["ФИО", "Отдел"])
    ws.append(["Сидоров С. С.", "Аналитика"])

    saved = {}
    monkeypatch.setattr("src.logic.read_sprav_file", lambda _: sprav)
    monkeypatch.setattr("src.logic.load_check_file", lambda *a, **k: (wb, ws, None))
    monkeypatch.setattr(
        "src.logic.save_file",
        lambda workbook, path: saved.setdefault("files", {}).__setitem__(str(path), workbook),
    )

    from src.logic import MSG_DUP_ALL_DISMISSED

    results, stats = run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=True,
        check_duplicates=False,
    )
    assert (ws["E2"].value or "") == MSG_DUP_ALL_DISMISSED
    assert is_accepted_for_final(MSG_DUP_ALL_DISMISSED) is False
    assert stats["final_rows"] == 0
    clean = saved["files"]["out.xlsx"].active
    assert clean["A2"].value is None


def test_include_unresolved_duplicates_in_final(monkeypatch):
    sprav = pd.DataFrame(
        [
            {
                "ФИО": "Смирнов Алексей Николаевич",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
            {
                "ФИО": "Смирнов Алексей Николаевич",
                "Штатная должность": "Программист",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
        ]
    )
    wb = Workbook()
    ws = wb.active
    ws.append(["ФИО", "Отдел"])
    ws.append(["Смирнов А. Н.", "IT"])

    saved = {}
    monkeypatch.setattr("src.logic.read_sprav_file", lambda _: sprav)
    monkeypatch.setattr("src.logic.load_check_file", lambda *a, **k: (wb, ws, None))
    monkeypatch.setattr(
        "src.logic.save_file",
        lambda workbook, path: saved.setdefault("files", {}).__setitem__(str(path), workbook),
    )

    _results, stats_off = run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out_off.xlsx",
        check_duplicates=True,
        prompt_duplicate_selection=False,
        include_unresolved_duplicates=False,
    )
    assert saved["files"]["out_off.xlsx"].active["A2"].value is None

    _results, stats_on = run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out_on.xlsx",
        check_duplicates=True,
        prompt_duplicate_selection=False,
        include_unresolved_duplicates=True,
    )
    assert saved["files"]["out_on.xlsx"].active["A2"].value == "Смирнов А. Н."
    assert stats_on["final_rows"] == 1
    assert stats_off["final_rows"] == 0


def test_pick_default_prefers_active():
    records = [
        {
            "ФИО": "A",
            "Штатная должность": "Интеграция: штатная должность по умолч.",
            "Статус занятости": "уволился",
            "Увольн.": "01.01.2020",
        },
        {
            "ФИО": "A",
            "Штатная должность": "Инженер",
            "Статус занятости": "работает",
            "Увольн.": "31.12.9999",
        },
    ]
    chosen = pick_default_duplicate_record(records)
    assert chosen["Штатная должность"] == "Инженер"


def test_pick_default_returns_none_when_all_inactive():
    records = [
        {
            "ФИО": "A",
            "Штатная должность": "Интеграция: штатная должность по умолч.",
            "Статус занятости": "уволился",
            "Увольн.": "01.01.2020",
        },
        {
            "ФИО": "A",
            "Группа сотрудников": "Внутренние совмест.",
            "Штатная должность": "Инженер",
            "Статус занятости": "работает",
            "Увольн.": "31.12.9999",
        },
    ]
    assert pick_default_duplicate_record(records) is None
    assert pick_default_duplicate_record([records[0]]) is None

