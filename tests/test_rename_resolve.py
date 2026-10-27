"""Тесты авторазрешения смены ФИО."""

import pandas as pd
from openpyxl import Workbook

from src.check_change_fio import change_fio, chainge_fio, prepare_rename_df, resolve_via_rename
from src.constants import ELOYY_FILL_COLOR
from src.logic import MSG_RENAMED_FOUND, run_check


def _sprav_df():
    return pd.DataFrame(
        [
            {
                "ФИО": "Кузнецова Анна Сергеевна",
                "Штатная должность": "Бухгалтер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
                "Таб.№": "2003",
                "Подразделение": "Бухгалтерия",
            },
            {
                "ФИО": "Смирнов Алексей Николаевич",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
                "Таб.№": "2018",
                "Подразделение": "IT",
            },
            {
                "ФИО": "Смирнов Алексей Николаевич",
                "Штатная должность": "Программист",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
                "Таб.№": "2018",
                "Подразделение": "IT",
            },
        ]
    )


def test_resolve_via_rename_old_to_new(tmp_path):
    path = tmp_path / "renamed.xlsx"
    pd.DataFrame(
        {
            "Старое ФИО": ["Иванова Анна Сергеевна"],
            "Новое ФИО": ["Кузнецова Анна Сергеевна"],
            "Таб.№": ["2003"],
        }
    ).to_excel(path, index=False)

    rename_df = prepare_rename_df(str(path))
    assert rename_df is not None

    from src.fio_index import build_fio_index

    fio_index = build_fio_index(_sprav_df())
    records, via = resolve_via_rename("Иванова А. С.", rename_df, fio_index)
    assert records is not None
    assert len(records) == 1
    assert records[0]["ФИО"] == "Кузнецова Анна Сергеевна"
    assert via is not None


def test_rename_same_person_duplicate_records(tmp_path):
    """Смена ФИО на человека с несколькими записями одного person_key — resolve ок."""
    path = tmp_path / "renamed.xlsx"
    pd.DataFrame(
        {
            "Старое ФИО": ["Иванов Алексей Николаевич"],
            "Новое ФИО": ["Смирнов Алексей Николаевич"],
            "Таб.№": ["9999"],
        }
    ).to_excel(path, index=False)

    rename_df = prepare_rename_df(str(path))
    from src.fio_index import build_fio_index

    fio_index = build_fio_index(_sprav_df())
    records, via = resolve_via_rename("Иванов А. Н.", rename_df, fio_index)
    assert records is not None
    assert len(records) == 2
    assert via is not None


def test_rename_ambiguous_two_people(tmp_path):
    """Две разные цели в файле смены → None."""
    path = tmp_path / "renamed.xlsx"
    pd.DataFrame(
        {
            "Старое ФИО": ["Старое Имя Имя", "Старое Имя Имя"],
            "Новое ФИО": ["Кузнецова Анна Сергеевна", "Смирнов Алексей Николаевич"],
            "Таб.№": ["2003", "2018"],
        }
    ).to_excel(path, index=False)

    rename_df = prepare_rename_df(str(path))
    from src.fio_index import build_fio_index

    sprav = _sprav_df()
    fio_index = build_fio_index(sprav)
    tab_index = {"2003": [sprav.iloc[0]], "2018": [sprav.iloc[1]]}
    records, via = resolve_via_rename("Старое Имя Имя", rename_df, fio_index, tab_index)
    assert records is None
    assert via is None


def test_chainge_fio_reuses_rename_df_without_reread(tmp_path, monkeypatch):
    path = tmp_path / "renamed.xlsx"
    pd.DataFrame(
        {
            "Старое ФИО": ["Неизвестный Николай Николаевич"],
            "Новое ФИО": ["Кузнецова Анна Сергеевна"],
            "Таб.№": ["2003"],
        }
    ).to_excel(path, index=False)

    rename_df = prepare_rename_df(str(path))
    calls = {"n": 0}

    def boom(*_a, **_k):
        calls["n"] += 1
        raise AssertionError("should not re-read rename file")

    monkeypatch.setattr("src.check_change_fio.read_sprav_file", boom)
    monkeypatch.setattr("src.check_change_fio.prepare_rename_df", boom)

    from src.fio_index import build_fio_index

    logs = []
    sprav = _sprav_df()
    chainge_fio(
        ["Неизвестный Н. Н."],
        chainge_fio_file=str(path),
        write=logs.append,
        fio_index=build_fio_index(sprav),
        rename_df=rename_df,
        tab_index={"2003": [sprav.iloc[0]]},
    )
    assert calls["n"] == 0
    assert any("Кузнецова" in line for line in logs)


def test_run_check_applies_rename_into_final(monkeypatch, tmp_path):
    sprav = _sprav_df().iloc[:2].copy()
    rename_path = tmp_path / "renamed.xlsx"
    pd.DataFrame(
        {
            "Старое ФИО": ["Иванова Анна Сергеевна"],
            "Новое ФИО": ["Кузнецова Анна Сергеевна"],
            "Таб.№": ["2003"],
        }
    ).to_excel(rename_path, index=False)

    wb = Workbook()
    ws = wb.active
    ws.append(["ФИО", "Отдел"])
    ws.append(["Иванова А. С.", "Бухгалтерия"])
    ws.append(["Неизвестный Н. Н.", "HR"])

    saved = {}
    monkeypatch.setattr(
        "src.logic.read_sprav_file",
        lambda path: sprav if "sprav" in str(path) else pd.read_excel(path, dtype=str).fillna(""),
    )
    monkeypatch.setattr(
        "src.check_change_fio.read_sprav_file",
        lambda path: pd.read_excel(path, dtype=str).fillna(""),
    )
    monkeypatch.setattr("src.logic.load_check_file", lambda *a, **k: (wb, ws, None))
    monkeypatch.setattr(
        "src.logic.save_file",
        lambda workbook, path: saved.setdefault("files", {}).__setitem__(str(path), workbook),
    )

    results = run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=True,
        check_duplicates=False,
        renamed_fio_file=str(rename_path),
        parent=None,
    )
    if isinstance(results, tuple):
        results = results[0]

    assert (ws["E2"].value or "") == MSG_RENAMED_FOUND
    assert "Не найден" in (ws["E3"].value or "")
    clean = saved["files"]["out.xlsx"].active
    assert clean["A2"].value == "Иванова А. С."
    assert clean["A3"].value is None
    assert any(r.get("status") == "green" and "смены ФИО" in r.get("comment", "") for r in results)


def test_run_check_rename_plus_duplicate_paint_dedicated_cols(monkeypatch, tmp_path):
    """Rename → дубликат: подсветка только в служебных колонках дублей, не в «Отдел»."""
    sprav = _sprav_df()
    rename_path = tmp_path / "renamed.xlsx"
    pd.DataFrame(
        {
            "Старое ФИО": ["Иванов Алексей Николаевич"],
            "Новое ФИО": ["Смирнов Алексей Николаевич"],
            "Таб.№": ["2018"],
        }
    ).to_excel(rename_path, index=False)

    wb = Workbook()
    ws = wb.active
    ws.append(["ФИО", "Отдел"])
    ws.append(["Иванов А. Н.", "IT"])

    monkeypatch.setattr(
        "src.logic.read_sprav_file",
        lambda path: sprav if "sprav" in str(path) else pd.read_excel(path, dtype=str).fillna(""),
    )
    monkeypatch.setattr(
        "src.check_change_fio.read_sprav_file",
        lambda path: pd.read_excel(path, dtype=str).fillna(""),
    )
    monkeypatch.setattr("src.logic.load_check_file", lambda *a, **k: (wb, ws, None))
    monkeypatch.setattr("src.logic.save_file", lambda *_: None)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_duplicates=True,
        prompt_duplicate_selection=False,
        renamed_fio_file=str(rename_path),
    )

    def _col(header: str) -> int:
        for col in range(1, (ws.max_column or 1) + 1):
            if ws.cell(row=1, column=col).value == header:
                return col
        raise AssertionError(header)

    assert ELOYY_FILL_COLOR in (ws["A2"].fill.start_color.rgb or "")
    # Регрессия: пользовательская колонка «Отдел» не красится шагами 2/3
    assert ELOYY_FILL_COLOR not in (ws["B2"].fill.start_color.rgb or "")
    step2 = _col("Дубль шаг 2")
    step3 = _col("Дубль шаг 3")
    # Два активных дубля одного ФИО — остаются жёлтыми в служебных колонках
    assert ELOYY_FILL_COLOR in (ws.cell(row=2, column=step2).fill.start_color.rgb or "")
    assert ELOYY_FILL_COLOR in (ws.cell(row=2, column=step3).fill.start_color.rgb or "")
    assert _col("Дубль выбор")
