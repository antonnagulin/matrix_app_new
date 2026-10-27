"""Тесты src/write_position_hierarchy.py."""

from openpyxl import Workbook

from src.write_position_hierarchy import write_position_hierarchy


def test_writes_hierarchy_and_position():
    wb = Workbook()
    ws = wb.active
    not_d = {
        (5, "x", "Иванов", "IT"): [
            {
                "Орг. единица уровня 1": "Головной офис",
                "Орг. единица уровня 2": "IT",
                "Орг. единица уровня 3": "",
                "Орг. единица уровня 4": "",
                "Штатная должность": "Инженер",
            }
        ]
    }
    write_position_hierarchy(ws, not_d, new_col_hierarchy_index=3, new_col_position_index=4)
    assert ws.cell(row=5, column=3).value == "Головной офис / IT"
    assert ws.cell(row=5, column=4).value == "Инженер"
