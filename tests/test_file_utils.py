"""Тесты file_utils — чтение/запись Excel и заливки."""

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from src.file_utils import (
    read_sprav_file,
    load_check_file,
    save_file,
    get_fill,
    clear_fill,
    find_header_row,
    peek_sheet_rows,
    row_values,
)
from src.constants import RED_FILL_COLOR


def test_read_sprav_fills_na_with_empty_string(tmp_path):
    path = tmp_path / "sprav.xlsx"
    pd.DataFrame({"ФИО": ["А", None], "Увольн.": [None, "01.01.2020"]}).to_excel(
        path, index=False
    )

    df = read_sprav_file(path)

    assert df.loc[1, "ФИО"] == ""
    assert df.loc[0, "Увольн."] == ""


def test_load_and_save_roundtrip(tmp_path):
    src = tmp_path / "in.xlsx"
    dst = tmp_path / "out.xlsx"

    wb = Workbook()
    wb.active["A1"] = "ФИО"
    wb.active["A2"] = "Тестов Т. Т."
    wb.save(src)

    loaded_wb, loaded_ws, values_ws = load_check_file(src)
    assert loaded_ws["A2"].value == "Тестов Т. Т."
    assert values_ws is not None

    save_file(loaded_wb, dst)
    assert load_workbook(dst).active["A2"].value == "Тестов Т. Т."


def test_get_fill_returns_solid_pattern():
    fill = get_fill(RED_FILL_COLOR)
    assert fill.fill_type == "solid"
    assert fill.start_color.rgb.endswith(RED_FILL_COLOR)


def test_find_header_row_on_first_line():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "ФИО"
    ws["B1"] = "Отдел"
    ws["A2"] = "Иванов И. И."

    assert find_header_row(ws) == 1


def test_find_header_row_shifted():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Отчёт по сотрудникам"
    ws["A2"] = "Период: 2026"
    ws["A12"] = "ФИО"
    ws["B12"] = "Подразделение"
    ws["A13"] = "Петров П. П."

    assert find_header_row(ws) == 12


def test_find_header_row_missing_raises():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Название"
    ws["A2"] = "Данные"

    with pytest.raises(ValueError, match="Не найдена строка заголовков"):
        find_header_row(ws, max_scan=5)


def test_peek_sheet_rows(tmp_path):
    path = tmp_path / "check.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A5"] = "ФИО"
    ws["A6"] = "Сидоров С. С."
    wb.save(path)

    rows, detected = peek_sheet_rows(path, max_rows=10)
    assert detected == 5
    assert rows[4][0] == 5
    assert "ФИО" in rows[4]


def test_clear_fill():
    fill = clear_fill()
    assert fill.fill_type is None


def test_row_values():
    wb = Workbook()
    ws = wb.active
    ws["A3"] = "ФИО"
    ws["B3"] = "Отдел"
    ws["C3"] = None

    assert row_values(ws, 3) == ["ФИО", "Отдел", None]


def test_list_sheet_names_and_load_named_sheet(tmp_path):
    from src.file_utils import list_sheet_names, cell_value

    path = tmp_path / "multi.xlsx"
    wb = Workbook()
    wb.active.title = "Первый"
    wb.active["A1"] = "ФИО"
    second = wb.create_sheet("Второй")
    second["A1"] = "ФИО"
    second["A2"] = "С листа 2"
    wb.save(path)

    assert list_sheet_names(path) == ["Первый", "Второй"]
    _wb, ws, _values = load_check_file(path, sheet_name="Второй")
    assert ws["A2"].value == "С листа 2"


def test_cell_value_merged_and_formula_cache(tmp_path):
    from src.file_utils import cell_value

    path = tmp_path / "merge.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "ФИО"
    ws.merge_cells("A2:B2")
    ws["A2"] = "Иванов И. И."
    wb.save(path)

    loaded_wb, loaded_ws, values_ws = load_check_file(path)
    assert cell_value(loaded_ws, 2, 1, values_ws) == "Иванов И. И."
    assert cell_value(loaded_ws, 2, 2, values_ws) == "Иванов И. И."


def test_save_file_permission_error_message(tmp_path, monkeypatch):
    wb = Workbook()
    path = tmp_path / "locked.xlsx"

    def boom(*_a, **_k):
        raise PermissionError(13, "denied")

    monkeypatch.setattr(wb, "save", boom)
    with pytest.raises(PermissionError, match="откройте|закройте|Excel"):
        save_file(wb, path)


def test_assert_safe_xlsx_rejects_empty(tmp_path):
    from src.file_utils import assert_safe_xlsx_path

    path = tmp_path / "empty.xlsx"
    path.write_bytes(b"")
    with pytest.raises(ValueError, match="пустой"):
        assert_safe_xlsx_path(path)


def test_assert_safe_xlsx_rejects_huge(tmp_path, monkeypatch):
    from src.file_utils import assert_safe_xlsx_path

    path = tmp_path / "huge.xlsx"
    path.write_bytes(b"PK")  # minimal non-empty

    class Stat:
        st_size = 300 * 1024 * 1024

    monkeypatch.setattr(type(path), "stat", lambda self: Stat())
    with pytest.raises(ValueError, match="слишком большой"):
        assert_safe_xlsx_path(path)


def test_load_check_rejects_corrupt(tmp_path):
    path = tmp_path / "bad.xlsx"
    path.write_bytes(b"not-an-excel-file")
    with pytest.raises(ValueError, match="повреждён|не Excel"):
        load_check_file(path)


def test_iter_column_values_fast_path():
    from src.file_utils import iter_column_values

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "ФИО"
    ws["A2"] = "Иванов"
    ws["A3"] = "Петров"
    values = list(iter_column_values(ws, 1, 2, 3))
    assert values == [(2, "Иванов"), (3, "Петров")]
