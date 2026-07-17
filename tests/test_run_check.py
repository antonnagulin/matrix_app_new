from openpyxl import Workbook
import pandas as pd
import openpyxl

from src.logic import run_check, SKIP_ALL_DUPLICATES
from src.constants import RED_FILL_COLOR, ORANGE_FILL_COLOR, ELOYY_FILL_COLOR, GREEN_FILL_COLOR


class DummyLog:
    def __init__(self):
        self.lines = []

    def insert(self, *_args):
        self.lines.append(_args[1])

    def see(self, *_):
        pass

    def text(self):
        return "".join(self.lines)


def _make_check_wb(rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["ФИО", "Отдел"])
    for fio, dept in rows:
        ws.append([fio, dept])
    return wb, ws


def _patch_io(monkeypatch, sprav_df, wb, ws, saved=None):
    monkeypatch.setattr("src.logic.read_sprav_file", lambda _: sprav_df)
    monkeypatch.setattr("src.logic.load_check_file", lambda _: (wb, ws))

    def _save(workbook, path):
        if saved is not None:
            saved["path"] = path
            saved["wb"] = workbook

    monkeypatch.setattr("src.logic.save_file", _save)


def test_marks_missing_employee_red(monkeypatch):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            }
        ]
    )
    wb, ws = _make_check_wb(
        [
            ("Иванов Иван Иванович", "IT"),
            ("Петров Петр Петрович", "HR"),
        ]
    )
    log = DummyLog()
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check("sprav.xlsx", "check.xlsx", "out.xlsx", log=log)

    assert RED_FILL_COLOR in (ws["A3"].fill.start_color.rgb or "")
    assert "Петров Петр Петрович" in log.text()
    assert "отсутствующие в справочнике" in log.text()


def test_marks_dismissed_employee_orange(monkeypatch):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Сидоров Сидор Сидорович",
                "Штатная должность": "Интеграция: штатная должность по умолч.",
                "Статус занятости": "уволился",
                "Увольн.": "01.01.2020",
            }
        ]
    )
    wb, ws = _make_check_wb([("Сидоров Сидор Сидорович", "IT")])
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check("sprav.xlsx", "check.xlsx", "out.xlsx", check_dismissed=True)

    assert ORANGE_FILL_COLOR in (ws["A2"].fill.start_color.rgb or "")


def test_skips_dismissed_check_when_disabled(monkeypatch):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Сидоров Сидор Сидорович",
                "Штатная должность": "Интеграция: штатная должность по умолч.",
                "Статус занятости": "уволился",
                "Увольн.": "01.01.2020",
            }
        ]
    )
    wb, ws = _make_check_wb([("Сидоров Сидор Сидорович", "IT")])
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=False,
        check_duplicates=False,
    )

    fill = ws["A2"].fill.start_color.rgb or "00000000"
    assert ORANGE_FILL_COLOR not in fill
    assert RED_FILL_COLOR not in fill


def test_skips_empty_fio_rows(monkeypatch):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            }
        ]
    )
    wb, ws = _make_check_wb(
        [
            ("", "IT"),
            ("Иванов Иван Иванович", "IT"),
        ]
    )
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=False,
        check_duplicates=False,
    )

    fill = ws["A2"].fill.start_color.rgb or "00000000"
    assert RED_FILL_COLOR not in fill


def test_asks_once_for_duplicate_missing_fio(monkeypatch):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            }
        ]
    )
    wb, ws = _make_check_wb(
        [
            ("Неизвестный Н. Н.", "IT"),
            ("Неизвестный Н.Н.", "HR"),
        ]
    )
    calls = {"count": 0}

    def fake_ask(*_args, **_kwargs):
        calls["count"] += 1
        return "no"

    monkeypatch.setattr("src.logic.ask_include_missing", fake_ask)
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=False,
        check_duplicates=False,
        parent=object(),
    )

    # Диалог и красная подсветка — только для первого вхождения ФИО
    assert calls["count"] == 1
    assert RED_FILL_COLOR in (ws["A2"].fill.start_color.rgb or "")


def test_repeat_missing_fio_does_not_crash_with_dismissed_check(monkeypatch):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            }
        ]
    )
    wb, ws = _make_check_wb(
        [
            ("Неизвестный Н. Н.", "IT"),
            ("Неизвестный Н.Н.", "HR"),
        ]
    )
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=True,
        check_duplicates=False,
        parent=None,
    )


def test_creates_included_report_when_user_confirms(monkeypatch, tmp_path):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            }
        ]
    )
    wb, ws = _make_check_wb([("Новиков Н. Н.", "IT")])
    output = tmp_path / "result.xlsx"
    saved = {}

    monkeypatch.setattr("src.logic.ask_include_missing", lambda *_a, **_k: "yes")
    _patch_io(monkeypatch, sprav_df, wb, ws, saved=saved)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        str(output),
        check_dismissed=False,
        check_duplicates=False,
        parent=object(),
    )

    included = tmp_path / "result_included.xlsx"
    assert included.exists()
    report = openpyxl.load_workbook(included)
    assert report.active["A1"].value == "ФИО"
    assert report.active["A2"].value == "Новиков Н. Н."
    assert report.active["B2"].value == "IT"


def test_yes_all_includes_subsequent_missing_without_prompt(monkeypatch, tmp_path):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            }
        ]
    )
    wb, ws = _make_check_wb(
        [
            ("Первый Н. Н.", "IT"),
            ("Второй В. В.", "HR"),
            ("Третий Т. Т.", "Склад"),
        ]
    )
    output = tmp_path / "result.xlsx"
    calls = {"count": 0}

    def fake_ask(_parent, fio):
        calls["count"] += 1
        return "yes_all" if fio == "Первый Н. Н." else "no"

    monkeypatch.setattr("src.logic.ask_include_missing", fake_ask)
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        str(output),
        check_dismissed=False,
        check_duplicates=False,
        parent=object(),
    )

    assert calls["count"] == 1
    included = tmp_path / "result_included.xlsx"
    assert included.exists()
    report = openpyxl.load_workbook(included)
    assert report.active["A2"].value == "Первый Н. Н."
    assert report.active["A3"].value == "Второй В. В."
    assert report.active["A4"].value == "Третий Т. Т."


def test_duplicates_logged_and_selection_called(monkeypatch):
    # Два активных дубликата — остаются неоднозначными, нужен выбор
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Менеджер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
        ]
    )
    wb, ws = _make_check_wb([("Иванов Иван Иванович", "IT")])
    log = DummyLog()
    selected = {"called": False}

    def fake_choose(parent, fio, records):
        selected["called"] = True
        selected["fio"] = fio
        selected["count"] = len(records)
        return records[0]

    monkeypatch.setattr("src.logic.choose_actual_record", fake_choose)
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=True,
        check_duplicates=True,
        log=log,
        parent=object(),
    )

    assert selected["called"] is True
    assert selected["count"] == 2
    assert "Дубликаты в справочнике для" in log.text()
    assert "Выбрана актуальная запись" in log.text()
    assert GREEN_FILL_COLOR in (ws["C2"].fill.start_color.rgb or "")
    assert ELOYY_FILL_COLOR in (ws["A2"].fill.start_color.rgb or "")


def test_skip_all_duplicates_skips_remaining_prompts(monkeypatch):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Менеджер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
            {
                "ФИО": "Петров Петр Петрович",
                "Штатная должность": "Аналитик",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
            {
                "ФИО": "Петров Петр Петрович",
                "Штатная должность": "Юрист",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
        ]
    )
    wb, ws = _make_check_wb(
        [
            ("Иванов Иван Иванович", "IT"),
            ("Петров Петр Петрович", "HR"),
        ]
    )
    log = DummyLog()
    calls = {"count": 0}

    def fake_choose(parent, fio, records):
        calls["count"] += 1
        return SKIP_ALL_DUPLICATES

    monkeypatch.setattr("src.logic.choose_actual_record", fake_choose)
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=True,
        check_duplicates=True,
        log=log,
        parent=object(),
    )

    assert calls["count"] == 1
    assert "Дальнейший выбор дубликатов пропущен" in log.text()


def test_duplicates_logged_without_prompt_when_disabled(monkeypatch):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Менеджер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
        ]
    )
    wb, ws = _make_check_wb([("Иванов Иван Иванович", "IT")])
    log = DummyLog()
    calls = {"count": 0}

    def fake_choose(*_args, **_kwargs):
        calls["count"] += 1
        return None

    monkeypatch.setattr("src.logic.choose_actual_record", fake_choose)
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=True,
        check_duplicates=True,
        prompt_duplicate_selection=False,
        log=log,
        parent=object(),
    )

    assert calls["count"] == 0
    assert "Дубликаты в справочнике для" in log.text()
    assert "Дубликаты записаны в журнал без ручного выбора." in log.text()
    assert ELOYY_FILL_COLOR in (ws["A2"].fill.start_color.rgb or "")


def test_no_duplicates_message(monkeypatch):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Интеграция: штатная должность по умолч.",
                "Статус занятости": "уволился",
                "Увольн.": "01.01.2020",
            }
        ]
    )
    wb, ws = _make_check_wb([("Иванов Иван Иванович", "IT")])
    log = DummyLog()
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=True,
        check_duplicates=True,
        log=log,
    )

    assert "Дубликатов не найдено" in log.text()


def test_saves_output_file(monkeypatch):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            }
        ]
    )
    wb, ws = _make_check_wb([("Иванов Иван Иванович", "IT")])
    saved = {}
    _patch_io(monkeypatch, sprav_df, wb, ws, saved=saved)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "custom_out.xlsx",
        check_dismissed=False,
        check_duplicates=False,
    )

    assert saved["path"] == "custom_out.xlsx"
    assert saved["wb"] is wb


def test_no_answer_does_not_create_included_report(monkeypatch, tmp_path):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            }
        ]
    )
    wb, ws = _make_check_wb(
        [
            ("Фёдоров Ф. Ф.", "IT"),
            ("Неизвестный Н. Н.", "HR"),
        ]
    )
    output = tmp_path / "result.xlsx"
    included = tmp_path / "result_included.xlsx"
    included.write_bytes(b"stale")

    monkeypatch.setattr("src.logic.ask_include_missing", lambda *_a, **_k: "no")
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        str(output),
        check_dismissed=False,
        check_duplicates=False,
        parent=object(),
    )

    assert not included.exists()


def test_no_all_excludes_subsequent_missing(monkeypatch, tmp_path):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            }
        ]
    )
    wb, ws = _make_check_wb(
        [
            ("Первый Н. Н.", "IT"),
            ("Второй В. В.", "HR"),
        ]
    )
    output = tmp_path / "result.xlsx"
    calls = {"count": 0}

    def fake_ask(_parent, fio):
        calls["count"] += 1
        return "no_all" if fio == "Первый Н. Н." else "no"

    monkeypatch.setattr("src.logic.ask_include_missing", fake_ask)
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        str(output),
        check_dismissed=False,
        check_duplicates=False,
        parent=object(),
    )

    assert calls["count"] == 1
    assert not (tmp_path / "result_included.xlsx").exists()


def test_mixed_duplicate_resolves_to_green(monkeypatch):
    """Активная + уволенная запись → автоматически зелёный, без диалога."""
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Сидоров Сидор Сидорович",
                "Штатная должность": "Аналитик",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            },
            {
                "ФИО": "Сидоров Сидор Сидорович",
                "Штатная должность": "Интеграция: штатная должность по умолч.",
                "Статус занятости": "уволился",
                "Увольн.": "15.03.2023",
            },
        ]
    )
    wb, ws = _make_check_wb([("Сидоров С. С.", "IT")])
    calls = {"count": 0}

    def fake_choose(*_a, **_k):
        calls["count"] += 1
        return None

    monkeypatch.setattr("src.logic.choose_actual_record", fake_choose)
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=True,
        check_duplicates=True,
        parent=object(),
    )

    assert calls["count"] == 0
    assert GREEN_FILL_COLOR in (ws["B2"].fill.start_color.rgb or "")
    assert ELOYY_FILL_COLOR in (ws["A2"].fill.start_color.rgb or "")


def test_all_dismissed_duplicates_mark_orange(monkeypatch):
    """Все записи дубликата уволены → оранжевый."""
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Сидоров Сидор Сидорович",
                "Штатная должность": "Интеграция: штатная должность по умолч.",
                "Статус занятости": "уволился",
                "Увольн.": "15.03.2023",
            },
            {
                "ФИО": "Сидоров Сидор Сидорович",
                "Штатная должность": "Интеграция: штатная должность по умолч.",
                "Статус занятости": "уволился",
                "Увольн.": "01.01.2020",
            },
        ]
    )
    wb, ws = _make_check_wb([("Сидоров С. С.", "IT")])
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=True,
        check_duplicates=True,
        prompt_duplicate_selection=False,
    )

    assert ORANGE_FILL_COLOR in (ws["A2"].fill.start_color.rgb or "")


def test_active_employee_not_highlighted(monkeypatch):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            }
        ]
    )
    wb, ws = _make_check_wb([("Иванов И. И.", "IT")])
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=True,
        check_duplicates=False,
    )

    fill = ws["A2"].fill.start_color.rgb or "00000000"
    assert RED_FILL_COLOR not in fill
    assert ORANGE_FILL_COLOR not in fill
    assert ELOYY_FILL_COLOR not in fill


def test_header_row_auto_detect_shifted(monkeypatch):
    """Заголовки не в 1-й строке — автопоиск и проверка данных ниже."""
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
            }
        ]
    )
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Отчёт"
    ws["A2"] = "Период 2026"
    ws["A12"] = "ФИО"
    ws["B12"] = "Отдел"
    ws["A13"] = "Иванов И. И."
    ws["B13"] = "IT"
    ws["A14"] = "Петров П. П."
    ws["B14"] = "HR"
    log = DummyLog()
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=False,
        check_duplicates=False,
        log=log,
    )

    assert "Строка заголовков определена автоматически: 12" in log.text()
    assert RED_FILL_COLOR not in (ws["A13"].fill.start_color.rgb or "00000000")
    assert RED_FILL_COLOR in (ws["A14"].fill.start_color.rgb or "")


def test_header_row_manual_override(monkeypatch):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
                "Подразделение": "Отдел АСУ",
            }
        ]
    )
    wb = Workbook()
    ws = wb.active
    # Ложный «ФИО» в шапке + настоящие заголовки на 5-й
    ws["A1"] = "ФИО отчёта (не таблица)"
    ws["A5"] = "ФИО"
    ws["B5"] = "Отдел"
    ws["A6"] = "Неизвестный Н. Н."
    ws["B6"] = "HR"
    log = DummyLog()
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=False,
        check_duplicates=False,
        header_row=5,
        log=log,
    )

    assert "Строка заголовков задана вручную: 5" in log.text()
    assert RED_FILL_COLOR in (ws["A6"].fill.start_color.rgb or "")


def test_writes_department_and_position_from_sprav(monkeypatch):
    sprav_df = pd.DataFrame(
        [
            {
                "ФИО": "Иванов Иван Иванович",
                "Штатная должность": "Инженер",
                "Статус занятости": "работает",
                "Увольн.": "31.12.9999",
                "Подразделение": "Отдел АСУ",
            }
        ]
    )
    wb, ws = _make_check_wb([("Иванов И. И.", "IT")])
    _patch_io(monkeypatch, sprav_df, wb, ws)

    run_check(
        "sprav.xlsx",
        "check.xlsx",
        "out.xlsx",
        check_dismissed=False,
        check_duplicates=False,
    )

    assert ws["C1"].value == "Подразделение"
    assert ws["D1"].value == "Штатная должность (справочник)"
    assert ws["C2"].value == "Отдел АСУ"
    assert ws["D2"].value == "Инженер"
