"""Тесты src/paint_row_for_dict.py."""

from openpyxl import Workbook

from src.paint_row_for_dict import paint_rows_by_dict
from src.file_utils import get_fill
from src.constants import RED_FILL_COLOR, GREEN_FILL_COLOR


def _fill_rgb(cell):
    return cell.fill.start_color.rgb or ""


def test_paint_single_column():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "ФИО"
    ws["B1"] = "Отдел"
    ws["A2"] = "Иванов"
    ws["B2"] = "IT"
    fill = get_fill(RED_FILL_COLOR)
    data = {(2, "ИВАНОВ", "Иванов", "IT"): [{}]}

    paint_rows_by_dict(ws, data, fill, 1)

    assert RED_FILL_COLOR in _fill_rgb(ws["A2"])
    assert GREEN_FILL_COLOR not in _fill_rgb(ws["B2"])


def test_paint_multiple_columns():
    wb = Workbook()
    ws = wb.active
    ws["A2"] = "Иванов"
    ws["B2"] = "IT"
    fill = get_fill(GREEN_FILL_COLOR)
    data = {(2, "ИВАНОВ", "Иванов", "IT"): [{}]}

    paint_rows_by_dict(ws, data, fill, [1, 2])

    assert GREEN_FILL_COLOR in _fill_rgb(ws["A2"])
    assert GREEN_FILL_COLOR in _fill_rgb(ws["B2"])


def test_paint_entire_row():
    wb = Workbook()
    ws = wb.active
    ws["A2"] = "Иванов"
    ws["B2"] = "IT"
    ws["C2"] = "1001"
    fill = get_fill(RED_FILL_COLOR)
    data = {(2, "ИВАНОВ", "Иванов", "IT"): [{}]}

    paint_rows_by_dict(ws, data, fill)

    assert RED_FILL_COLOR in _fill_rgb(ws["A2"])
    assert RED_FILL_COLOR in _fill_rgb(ws["B2"])
    assert RED_FILL_COLOR in _fill_rgb(ws["C2"])
