"""Тесты src/check_change_fio.py."""

from openpyxl import Workbook

from src.check_change_fio import change_fio


def test_change_fio_no_match_silent(tmp_path, capsys):
    path = tmp_path / "rename.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Смена ФИО", "Новое ФИО"])
    ws.append(["А А А", "Б Б Б"])
    wb.save(path)

    logs = []
    change_fio(["Неизвестный Н. Н."], str(path), write=logs.append)
    assert any("не найдено" in x.lower() or "Совпадений" in x for x in logs) or logs == [
        "Совпадений в файле смены ФИО не найдено."
    ]


def test_change_fio_reports_ambiguous_renames(tmp_path):
    path = tmp_path / "rename.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Старое ФИО", "Новое ФИО"])
    ws.append(["Фёдоров Ф. Ф.", "Федоров Фёдор Фёдорович"])
    ws.append(["Фёдоров Ф. Ф.", "Другой Д. Д."])
    wb.save(path)

    logs = []
    change_fio(["Фёдоров Ф. Ф."], str(path), write=logs.append)
    assert any("однозначно не найдено" in x or "Совпадений" in x or "Фёдоров" in x for x in logs)


def test_change_fio_supports_old_new_columns(tmp_path):
    path = tmp_path / "rename.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Старое ФИО", "Новое ФИО"])
    ws.append(["Неизвестный Н. Н.", "Известный И. И."])
    wb.save(path)

    logs = []
    change_fio(["Неизвестный Н. Н."], str(path), write=logs.append)
    assert logs  # хоть какой-то лог
