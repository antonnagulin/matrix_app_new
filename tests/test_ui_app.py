"""Тесты UI без полного mainloop."""

import tkinter as tk
from pathlib import Path
from unittest.mock import patch

import pytest

from src.ui.app import EmployeeCheckApp


@pytest.fixture
def app():
    root = tk.Tk()
    root.withdraw()
    application = EmployeeCheckApp(root)
    yield application
    root.destroy()


def test_validate_inputs_requires_sprav(app):
    app.sprav.set("")
    app.check.set("x.xlsx")
    app.output.set("out.xlsx")
    with patch("src.ui.app.messagebox.showerror") as mock_err:
        assert app._validate_inputs() is False
        mock_err.assert_called_once()


def test_validate_inputs_requires_existing_files(app, tmp_path):
    sprav = tmp_path / "sprav.xlsx"
    check = tmp_path / "check.xlsx"
    sprav.write_bytes(b"x")
    check.write_bytes(b"x")
    app.sprav.set(str(sprav))
    app.check.set(str(check))
    app.output.set(str(tmp_path / "out.xlsx"))
    assert app._validate_inputs() is True


def test_validate_inputs_rejects_output_equal_check(app, tmp_path):
    sprav = tmp_path / "sprav.xlsx"
    check = tmp_path / "check.xlsx"
    sprav.write_bytes(b"x")
    check.write_bytes(b"x")
    app.sprav.set(str(sprav))
    app.check.set(str(check))
    app.output.set(str(check))
    with patch("src.ui.app.messagebox.showerror") as mock_err:
        assert app._validate_inputs() is False
        mock_err.assert_called_once()
        assert "совпадает" in mock_err.call_args[0][1]


def test_suggest_output_name(app, tmp_path):
    check = tmp_path / "report.xlsx"
    check.write_bytes(b"x")
    app.output.set("check_result.xlsx")
    app._suggest_output_name(str(check))
    assert app.output.get() == str(tmp_path / "report_result.xlsx")


def test_detect_header_updates_spin(app, tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "check.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A10"] = "ФИО"
    ws["B10"] = "Отдел"
    wb.save(path)

    app.auto_header.set(True)
    app._detect_header_for_check(str(path))
    assert app.header_row.get() == 10
    assert "10" in app.header_info.get()


def test_on_close_while_running_closes_dialog_bridge(app):
    app._running = True
    with patch("src.ui.app.messagebox.askyesno", return_value=True):
        with patch.object(app.root, "destroy") as mock_destroy:
            app._on_close()
            assert app._dialog_bridge._closed is True
            mock_destroy.assert_called_once()


def test_dialog_bridge_closed_returns_default():
    from src.ui.dialog_bridge import TkDialogBridge

    root = tk.Tk()
    root.withdraw()
    bridge = TkDialogBridge(root, timeout_sec=0.3)
    bridge.close()
    assert bridge.ask_include_missing("X") == "no"
    root.destroy()
