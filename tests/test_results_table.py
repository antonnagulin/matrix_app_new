"""Тесты таблицы результатов и маппинга цветов."""

import tkinter as tk
from unittest.mock import patch

import pytest
from openpyxl import Workbook

from src.logic import (
    MSG_DISMISSED,
    MSG_DUP_STEP2_GREEN,
    MSG_NOT_FOUND,
    STATUS_GREEN,
    STATUS_ORANGE,
    STATUS_RED,
    build_final_workbook,
    collect_highlighted_rows,
    is_accepted_for_final,
    status_from_comment,
)
from src.ui.app import EmployeeCheckApp


@pytest.fixture
def app():
    root = tk.Tk()
    root.withdraw()
    application = EmployeeCheckApp(root)
    application._run_async = False
    yield application
    root.destroy()


def test_status_from_comment_maps_known_messages():
    assert status_from_comment(MSG_NOT_FOUND) == STATUS_RED
    assert status_from_comment(MSG_DISMISSED) == STATUS_ORANGE
    assert status_from_comment(MSG_DUP_STEP2_GREEN) == STATUS_GREEN
    assert status_from_comment("") is None
    assert status_from_comment(None) is None


def test_is_accepted_for_final():
    assert is_accepted_for_final(None) is True
    assert is_accepted_for_final("") is True
    assert is_accepted_for_final(MSG_DUP_STEP2_GREEN) is True
    assert is_accepted_for_final(MSG_NOT_FOUND) is False
    assert is_accepted_for_final(MSG_DISMISSED) is False
    from src.logic import MSG_DUP_DEFAULT, MSG_DUP_STEP3_YELLOW

    assert is_accepted_for_final(MSG_DUP_DEFAULT) is True
    assert is_accepted_for_final(MSG_DUP_STEP3_YELLOW) is False
    assert is_accepted_for_final(MSG_DUP_STEP3_YELLOW, include_unresolved_duplicates=True) is True


def test_build_final_workbook_keeps_only_accepted_without_comment():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "ФИО"
    ws["B1"] = "Подразделение"
    ws["C1"] = "Должность"
    ws["D1"] = "Комментарий"
    ws["E1"] = "Дубль шаг 2"
    ws["A2"] = "Иванов"
    ws["B2"] = "IT"
    ws["C2"] = "Инженер"
    ws["D2"] = ""
    ws["A3"] = "Петров"
    ws["D3"] = MSG_NOT_FOUND
    ws["A4"] = "Сидоров"
    ws["B4"] = "HR"
    ws["D4"] = MSG_DUP_STEP2_GREEN

    final = build_final_workbook(
        ws,
        headers=["ФИО", "Подразделение", "Должность", "Комментарий", "Дубль шаг 2"],
        fio_col=1,
        comment_col=4,
        data_start=2,
    ).active
    assert final["A1"].value == "ФИО"
    assert final["B1"].value == "Подразделение"
    assert final["C1"].value == "Должность"
    assert final["D1"].value is None
    assert final["E1"].value is None
    assert final["A2"].value == "Иванов"
    assert final["A3"].value == "Сидоров"
    assert final["A4"].value is None


def test_collect_highlighted_rows_skips_ok_rows():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "ФИО"
    ws["B1"] = "Подразделение"
    ws["C1"] = "Должность"
    ws["D1"] = "Комментарий"

    ws["A2"] = "Иванов Иван"
    ws["D2"] = MSG_NOT_FOUND
    ws["A3"] = "Петров Пётр"
    ws["B3"] = "Отдел А"
    ws["C3"] = "Инженер"
    ws["D3"] = ""
    ws["A4"] = "Сидоров Сидор"
    ws["D4"] = MSG_DISMISSED

    rows = collect_highlighted_rows(
        ws,
        fio_col=1,
        hierarchy_col=2,
        position_col=3,
        comment_col=4,
        data_start=2,
    )
    assert len(rows) == 2
    assert rows[0]["fio"] == "Иванов Иван"
    assert rows[0]["status"] == STATUS_RED
    assert rows[1]["fio"] == "Сидоров Сидор"
    assert rows[1]["status"] == STATUS_ORANGE


def test_fill_results_applies_color_tags(app):
    app._fill_results(
        [
            {
                "row": 2,
                "fio": "Иванов",
                "hierarchy": "Отдел",
                "position": "Инженер",
                "comment": MSG_NOT_FOUND,
                "status": "red",
            },
            {
                "row": 3,
                "fio": "Петров",
                "hierarchy": "",
                "position": "",
                "comment": MSG_DUP_STEP2_GREEN,
                "status": "green",
            },
        ]
    )
    items = app.results.get_children()
    assert len(items) == 2
    assert app.results.item(items[0], "tags") == ("red",)
    assert app.results.item(items[1], "tags") == ("green",)
    assert app.results.item(items[0], "values")[1] == "Иванов"


def test_on_run_fills_results_table(app, tmp_path):
    sample = [
        {
            "row": 5,
            "fio": "Тестов Т",
            "hierarchy": "Цех",
            "position": "Слесарь",
            "comment": MSG_DISMISSED,
            "status": "orange",
        }
    ]
    sprav = tmp_path / "sprav.xlsx"
    check = tmp_path / "check.xlsx"
    out = tmp_path / "out.xlsx"
    sprav.write_bytes(b"x")
    check.write_bytes(b"x")
    app.sprav.set(str(sprav))
    app.check.set(str(check))
    app.output.set(str(out))

    with patch("src.ui.app.run_check", return_value=sample), patch(
        "src.ui.app.messagebox.showinfo"
    ):
        app._on_run()

    assert not app._running
    items = app.results.get_children()
    assert len(items) == 1
    assert app.results.item(items[0], "tags") == ("orange",)
    assert "Готово" in app.status.get()
